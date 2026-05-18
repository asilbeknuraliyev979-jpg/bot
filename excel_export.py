import os
import pandas as pd
from pathlib import Path
from datetime import datetime
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from config.config import Config

def export_test_results(results_data: list, test_code: str, subject: str, format_type: str) -> str:
    """
    Exports a list of result dictionaries to the requested format (excel, csv, pdf).
    Each dictionary in results_data must contain:
      - full_name, phone, class_name, telegram_id, test_code, subject, correct_count, wrong_count, percentage, submission_time
    Includes the 'Jami savollar' (Total count) column.
    Returns the absolute path to the generated file.
    """
    Config.ensure_dirs()
    
    # Add dynamic ranking
    sorted_results = sorted(
        results_data, 
        key=lambda x: (x.get("percentage", 0), -x.get("duration_spent", 999999)), 
        reverse=True
    )
    
    formatted_data = []
    for rank, r in enumerate(sorted_results, 1):
        formatted_data.append({
            "O'rin (Rank)": rank,
            "Ism va Familiya": r.get("full_name", ""),
            "Sinf": r.get("class_name", ""),
            "Test Kodi": r.get("test_code", test_code),
            "Fan": r.get("subject", subject),
            "Jami savollar": r.get("correct_count", 0) + r.get("wrong_count", 0),
            "To'g'ri javoblar": r.get("correct_count", 0),
            "Noto'g'ri javoblar": r.get("wrong_count", 0),
            "Natija (%)": f"{r.get('percentage', 0.0)}%",
            "Topshirilgan vaqt": r.get("submission_time", "")
        })

    df = pd.DataFrame(formatted_data)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = Config.BASE_DIR / "storage" / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    
    filename = f"natijalar_{test_code}_{timestamp}"
    
    if format_type.lower() == "excel":
        file_path = export_dir / f"{filename}.xlsx"
        
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        df.to_excel(writer, sheet_name='Natijalar', index=False)
        
        worksheet = writer.sheets['Natijalar']
        
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        navy_header_fill = PatternFill(start_color='0D2040', end_color='0D2040', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        
        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        alt_row_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')
        white_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        gold_row_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        
        data_font = Font(name='Arial', size=10)
        bold_font = Font(name='Arial', size=10, bold=True)
        
        for row_num in range(2, len(df) + 2):
            if row_num == 2:
                row_fill = gold_row_fill
                row_font = bold_font
            else:
                row_fill = alt_row_fill if (row_num % 2 == 0) else white_row_fill
                row_font = data_font
            
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = row_font
                
                col_name = df.columns[col_num - 1]
                if col_name in ["O'rin (Rank)", "Sinf", "Test Kodi", "Jami savollar", "To'g'ri javoblar", "Noto'g'ri javoblar", "Natija (%)", "Topshirilgan vaqt"]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        writer.close()
        return str(file_path)
        
    elif format_type.lower() == "csv":
        file_path = export_dir / f"{filename}.csv"
        df.to_csv(file_path, index=False, encoding="utf-8-sig")
        return str(file_path)
        
    elif format_type.lower() == "pdf":
        file_path = export_dir / f"{filename}.pdf"
        
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=landscape(letter),
            rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#0D2040'),
            alignment=1,
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            name='MetaStyle',
            fontName='Helvetica',
            fontSize=11,
            alignment=1,
            spaceAfter=20
        )
        
        story = []
        story.append(Paragraph(f"TEST NATIJALARI HISOBOTI", title_style))
        story.append(Paragraph(f"Fan: {subject} | Test kodi: {test_code} | Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        
        headers = ["O'rin", "Ism va Familiya", "Sinf", "Jami", "To'g'ri", "Noto'g'ri", "Natija", "Topshirilgan vaqt"]
        table_rows = [headers]
        
        for row in formatted_data:
            table_rows.append([
                str(row["O'rin (Rank)"]),
                row["Ism va Familiya"],
                row["Sinf"],
                str(row["Jami savollar"]),
                str(row["To'g'ri javoblar"]),
                str(row["Noto'g'ri javoblar"]),
                row["Natija (%)"],
                row["Topshirilgan vaqt"]
            ])
            
        col_widths = [40, 160, 50, 45, 50, 55, 52, 120]
        
        t = Table(table_rows, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0D2040')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('TOPPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#FFF2CC')),
            ('FONTNAME', (0,1), (-1,1), 'Helvetica-Bold'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0,2), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('FONTNAME', (0,2), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
        ]))
        
        story.append(t)
        doc.build(story)
        return str(file_path)
    else:
        raise ValueError(f"Unsupported format type: {format_type}")

def export_statistics_report(stats_data: dict, format_type: str) -> str:
    """
    Exports system statistics to Excel or PDF.
    """
    Config.ensure_dirs()
    
    formatted_data = [
        {"Ko'rsatkich (Metric)": "Jami ro'yxatdan o'tgan o'quvchilar", "Qiymat (Value)": f"{stats_data['total_users']} ta"},
        {"Ko'rsatkich (Metric)": "Aktiv (Test topshirgan) o'quvchilar", "Qiymat (Value)": f"{stats_data['active_users']} ta"},
        {"Ko'rsatkich (Metric)": "Jami topshirilgan testlar", "Qiymat (Value)": f"{stats_data['total_results']} ta"},
        {"Ko'rsatkich (Metric)": "Muvaffaqiyatli berilgan sertifikatlar", "Qiymat (Value)": f"{stats_data['certificates_count']} ta"},
        {"Ko'rsatkich (Metric)": "Tizimdagi eng yuqori natija", "Qiymat (Value)": f"{stats_data['highest_score']}%"},
        {"Ko'rsatkich (Metric)": "Tizimdagi eng past natija", "Qiymat (Value)": f"{stats_data['lowest_score']}%"},
        {"Ko'rsatkich (Metric)": "O'rtacha o'zlashtirish ko'rsatkichi", "Qiymat (Value)": f"{stats_data['average_score']}%"}
    ]
    
    df = pd.DataFrame(formatted_data)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = Config.BASE_DIR / "storage" / "exports"
    
    filename = f"tizim_statistikasi_{timestamp}"
    
    if format_type.lower() == "excel":
        file_path = export_dir / f"{filename}.xlsx"
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        df.to_excel(writer, sheet_name='Statistika', index=False)
        worksheet = writer.sheets['Statistika']
        
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        navy_header_fill = PatternFill(start_color='0D2040', end_color='0D2040', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        alt_row_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')
        white_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        
        for row_num in range(2, len(df) + 2):
            row_fill = alt_row_fill if (row_num % 2 == 0) else white_row_fill
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name='Arial', size=10)
                if col_num == 1:
                    cell.alignment = left_align
                else:
                    cell.alignment = center_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 6, 20)
            
        writer.close()
        return str(file_path)
        
    elif format_type.lower() == "pdf":
        file_path = export_dir / f"{filename}.pdf"
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=letter,
            rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#0D2040'),
            alignment=1,
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            name='MetaStyle',
            fontName='Helvetica',
            fontSize=11,
            alignment=1,
            spaceAfter=25
        )
        
        story = []
        story.append(Paragraph(f"TIZIM STATISTIKASI HISOBOTI", title_style))
        story.append(Paragraph(f"Eksport qilingan vaqt: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        
        headers = ["Ko'rsatkich (Metric)", "Qiymat (Value)"]
        table_rows = [headers]
        for row in formatted_data:
            table_rows.append([row["Ko'rsatkich (Metric)"], row["Qiymat (Value)"]])
            
        t = Table(table_rows, colWidths=[320, 180])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0D2040')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,0), 'CENTER'),
            ('ALIGN', (0,1), (0,-1), 'LEFT'),
            ('ALIGN', (1,1), (1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 11),
            ('BOTTOMPADDING', (0,0), (-1,0), 10),
            ('TOPPADDING', (0,0), (-1,0), 10),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 10),
            ('BOTTOMPADDING', (0,1), (-1,-1), 8),
            ('TOPPADDING', (0,1), (-1,-1), 8),
        ]))
        
        story.append(t)
        doc.build(story)
        return str(file_path)
    else:
        raise ValueError(f"Unsupported format type: {format_type}")

def export_users_list(users_data: list, format_type: str) -> str:
    """
    Exports all registered users list to Excel or PDF for administrator view.
    """
    Config.ensure_dirs()
    
    formatted_data = []
    for idx, u in enumerate(users_data, 1):
        formatted_data.append({
            "№": idx,
            "Ism va Familiya": u.get("full_name", ""),
            "Sinf": u.get("class_name", ""),
            "Telefon raqam": u.get("phone", ""),
            "Telegram ID": u.get("telegram_id", ""),
            "Ro'yxatdan o'tgan sana": u.get("registration_date", "")
        })
        
    df = pd.DataFrame(formatted_data)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = Config.BASE_DIR / "storage" / "exports"
    
    filename = f"foydalanuvchilar_ruyxati_{timestamp}"
    
    if format_type.lower() == "excel":
        file_path = export_dir / f"{filename}.xlsx"
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        df.to_excel(writer, sheet_name='Foydalanuvchilar', index=False)
        worksheet = writer.sheets['Foydalanuvchilar']
        
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        navy_header_fill = PatternFill(start_color='0D2040', end_color='0D2040', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        alt_row_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')
        white_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        
        for row_num in range(2, len(df) + 2):
            row_fill = alt_row_fill if (row_num % 2 == 0) else white_row_fill
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name='Arial', size=10)
                
                col_name = df.columns[col_num - 1]
                if col_name in ["№", "Sinf", "Telegram ID", "Ro'yxatdan o'tgan sana"]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        writer.close()
        return str(file_path)
        
    elif format_type.lower() == "pdf":
        file_path = export_dir / f"{filename}.pdf"
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=letter,
            rightMargin=30, leftMargin=30, topMargin=35, bottomMargin=30
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            fontName='Helvetica-Bold',
            fontSize=15,
            textColor=colors.HexColor('#0D2040'),
            alignment=1,
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            name='MetaStyle',
            fontName='Helvetica',
            fontSize=10,
            alignment=1,
            spaceAfter=20
        )
        
        story = []
        story.append(Paragraph(f"RO'YXATDAN O'TGAN O'QUVCHILAR RO'YXATI", title_style))
        story.append(Paragraph(f"Jami ro'yxatdan o'tganlar: {len(users_data)} ta | Chop etilgan: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        
        headers = ["№", "Ism va Familiya", "Sinf", "Telefon raqam", "Telegram ID"]
        table_rows = [headers]
        
        for row in formatted_data:
            table_rows.append([
                str(row["№"]),
                row["Ism va Familiya"],
                row["Sinf"],
                row["Telefon raqam"],
                str(row["Telegram ID"])
            ])
            
        col_widths = [30, 200, 50, 110, 110]
        
        t = Table(table_rows, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0D2040')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('ALIGN', (1,1), (1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('TOPPADDING', (0,0), (-1,0), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
        ]))
        
        story.append(t)
        doc.build(story)
        return str(file_path)
    else:
        raise ValueError(f"Unsupported format type: {format_type}")