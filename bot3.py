

==================================================
FILE: .\loader.py
==================================================

import logging
from aiogram import Bot, Dispatcher
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from config.config import Config
from database.db import Database

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=[
        logging.StreamHandler()
    ]
)
logger = logging.getLogger("EduMasterBot")

# Initialize Bot
bot = Bot(
    token=Config.BOT_TOKEN,
    default=DefaultBotProperties(parse_mode=ParseMode.HTML)
)

# Initialize Dispatcher with Memory Storage for FSM
storage = MemoryStorage()
dp = Dispatcher(storage=storage)

# Initialize Database helper
db = Database()


==================================================
FILE: .\main.py
==================================================

import asyncio
import logging
from aiogram import Bot
from aiogram.types import BotCommand, BotCommandScopeDefault

from loader import bot, dp, db
from handlers import router as master_router
from config.config import Config

logger = logging.getLogger("EduMasterBot.main")

async def setup_commands(bot: Bot):
    """
    Sets up basic commands in the Telegram interface.
    """
    commands = [
        BotCommand(command="start", description="Botni ishga tushirish / qayta yuklash"),
        BotCommand(command="admin", description="Admin panelga kirish (faqat adminlar)")
    ]
    await bot.set_my_commands(commands, scope=BotCommandScopeDefault())

async def main():
    logger.info("🚀 Edu Master Bot starting...")
    
    # Ensure config storage folders exist
    Config.ensure_dirs()
    
    # Register master router which includes admin and user routers
    dp.include_router(master_router)
    
    # Register bot commands
    await setup_commands(bot)
    
    # Pre-populate default admin from .env in DB if not already exists
    if Config.ADMIN_ID:
        db.add_admin(Config.ADMIN_ID)
        logger.info(f"Added default admin ID: {Config.ADMIN_ID} to DB.")
        
    logger.info("✅ Bot is running. Polling...")
    
    # Start polling
    # Close any existing webhook (to prevent conflict errors)
    await bot.delete_webhook(drop_pending_updates=True)
    await dp.start_polling(bot)

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except (KeyboardInterrupt, SystemExit):
        logger.info("👋 Bot stopped.")


==================================================
FILE: .\config\config.py
==================================================

import os
from pathlib import Path
from dotenv import load_dotenv

# Calculate base directory (absolute path to edu_master_bot)
BASE_DIR = Path(__file__).resolve().parent.parent

# Load .env file
load_dotenv(dotenv_path=BASE_DIR / ".env")

class Config:
    BASE_DIR = BASE_DIR
    BOT_TOKEN = os.getenv("BOT_TOKEN", "8234813786:AAG3tV0mlQdj8S2gWnF4RVp9zuieTHCwdY8")

    ADMIN_ID = int(os.getenv("ADMIN_ID", "8420258761"))
    CHANNEL_USERNAME = os.getenv("CHANNEL_USERNAME", "@boliqulov700")

    # Folders for storage
    PDF_DIR = BASE_DIR / "storage" / "pdfs"
    CERTIFICATE_DIR = BASE_DIR / "storage" / "certificates"

    @classmethod
    def ensure_dirs(cls):
        cls.PDF_DIR.mkdir(parents=True, exist_ok=True)
        cls.CERTIFICATE_DIR.mkdir(parents=True, exist_ok=True)

# Ensure folders exist
Config.ensure_dirs()


==================================================
FILE: .\database\db.py
==================================================

import sqlite3
import os
from datetime import datetime
from pathlib import Path

class Database:
    def __init__(self, db_path=None):
        if db_path is None:
            # Place database in the same directory as the project
            self.db_path = str(Path(__file__).resolve().parent.parent / "edu_master.db")
        else:
            self.db_path = db_path
        self._init_db()

    def _get_connection(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            
            # Users table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    telegram_id INTEGER PRIMARY KEY,
                    full_name TEXT NOT NULL,
                    phone TEXT NOT NULL,
                    class_name TEXT NOT NULL,
                    registration_date TEXT NOT NULL,
                    is_blocked INTEGER DEFAULT 0
                )
            """)
            
            # Tests table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS tests (
                    code TEXT PRIMARY KEY,
                    subject TEXT NOT NULL,
                    pdf_path TEXT,
                    correct_answers TEXT NOT NULL,
                    duration INTEGER NOT NULL,
                    start_time TEXT NOT NULL,
                    end_time TEXT NOT NULL,
                    is_active INTEGER DEFAULT 1
                )
            """)
            
            # Results table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS results (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    telegram_id INTEGER NOT NULL,
                    test_code TEXT NOT NULL,
                    correct_count INTEGER NOT NULL,
                    wrong_count INTEGER NOT NULL,
                    percentage REAL NOT NULL,
                    submission_time TEXT NOT NULL,
                    duration_spent INTEGER NOT NULL,
                    FOREIGN KEY (telegram_id) REFERENCES users(telegram_id),
                    FOREIGN KEY (test_code) REFERENCES tests(code)
                )
            """)
            
            # Certificates table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS certificates (
                    id TEXT PRIMARY KEY,
                    telegram_id INTEGER NOT NULL,
                    test_code TEXT NOT NULL,
                    issue_date TEXT NOT NULL,
                    pdf_path TEXT,
                    image_path TEXT,
                    FOREIGN KEY (telegram_id) REFERENCES users(telegram_id),
                    FOREIGN KEY (test_code) REFERENCES tests(code)
                )
            """)
            
            # Admins table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS admins (
                    telegram_id INTEGER PRIMARY KEY,
                    added_date TEXT NOT NULL
                )
            """)
            
            # Channels table for mandatory subscription
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS channels (
                    channel_id TEXT PRIMARY KEY,
                    url TEXT NOT NULL
                )
            """)

            # Migration: Add is_gift to tests if not exists
            try:
                cursor.execute("ALTER TABLE tests ADD COLUMN is_gift INTEGER DEFAULT 0")
            except sqlite3.OperationalError:
                pass  # Already exists
                
            conn.commit()

    # User operations
    def get_user(self, telegram_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM users WHERE telegram_id = ?", (telegram_id,))
            row = cursor.fetchone()
            return dict(row) if row else None

    def add_user(self, telegram_id, full_name, phone, class_name):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            reg_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor.execute("""
                INSERT OR REPLACE INTO users (telegram_id, full_name, phone, class_name, registration_date, is_blocked)
                VALUES (?, ?, ?, ?, ?, 0)
            """, (telegram_id, full_name, phone, class_name, reg_date))
            conn.commit()

    def update_user(self, telegram_id, **kwargs):
        if not kwargs:
            return
        fields = ", ".join([f"{k} = ?" for k in kwargs.keys()])
        values = list(kwargs.values()) + [telegram_id]
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(f"UPDATE users SET {fields} WHERE telegram_id = ?", tuple(values))
            conn.commit()

    def get_all_users(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM users")
            return [dict(row) for row in cursor.fetchall()]

    def count_users(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*) FROM users")
            return cursor.fetchone()[0]

    # Admin operations
    def is_admin(self, telegram_id, default_admin_id=None):
        if default_admin_id and telegram_id == default_admin_id:
            return True
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT 1 FROM admins WHERE telegram_id = ?", (telegram_id,))
            return cursor.fetchone() is not None

    def add_admin(self, telegram_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            added_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor.execute("""
                INSERT OR IGNORE INTO admins (telegram_id, added_date)
                VALUES (?, ?)
            """, (telegram_id, added_date))
            conn.commit()

    def remove_admin(self, telegram_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM admins WHERE telegram_id = ?", (telegram_id,))
            conn.commit()

    def get_all_admins(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM admins")
            return [dict(row) for row in cursor.fetchall()]

    # Channels operations
    def add_channel(self, channel_id, url):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT OR REPLACE INTO channels (channel_id, url)
                VALUES (?, ?)
            """, (channel_id, url))
            conn.commit()

    def remove_channel(self, channel_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM channels WHERE channel_id = ?", (channel_id,))
            conn.commit()

    def get_all_channels(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM channels")
            return [dict(row) for row in cursor.fetchall()]

    # Test operations
    def get_test(self, code):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM tests WHERE code = ?", (code,))
            row = cursor.fetchone()
            return dict(row) if row else None

    def add_test(self, code, subject, pdf_path, correct_answers, duration, start_time, end_time, is_gift=0):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT OR REPLACE INTO tests (code, subject, pdf_path, correct_answers, duration, start_time, end_time, is_active, is_gift)
                VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)
            """, (code, subject, pdf_path, correct_answers, duration, start_time, end_time, is_gift))
            conn.commit()

    def get_gift_tests(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM tests WHERE is_gift = 1 ORDER BY code")
            return [dict(row) for row in cursor.fetchall()]

    def delete_test(self, code):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM tests WHERE code = ?", (code,))
            cursor.execute("DELETE FROM results WHERE test_code = ?", (code,))
            cursor.execute("DELETE FROM certificates WHERE test_code = ?", (code,))
            conn.commit()

    def get_all_tests(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM tests ORDER BY code")
            return [dict(row) for row in cursor.fetchall()]

    def update_test_status(self, code, is_active):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("UPDATE tests SET is_active = ? WHERE code = ?", (is_active, code))
            conn.commit()

    def update_test(self, code, **kwargs):
        if not kwargs:
            return
        fields = ", ".join([f"{k} = ?" for k in kwargs.keys()])
        values = list(kwargs.values()) + [code]
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(f"UPDATE tests SET {fields} WHERE code = ?", tuple(values))
            conn.commit()

    # Result operations
    def add_result(self, telegram_id, test_code, correct_count, wrong_count, percentage, duration_spent):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            sub_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor.execute("""
                INSERT INTO results (telegram_id, test_code, correct_count, wrong_count, percentage, submission_time, duration_spent)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, (telegram_id, test_code, correct_count, wrong_count, percentage, sub_time, duration_spent))
            conn.commit()
            return cursor.lastrowid

    def get_user_results(self, telegram_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT r.*, t.subject 
                FROM results r 
                JOIN tests t ON r.test_code = t.code 
                WHERE r.telegram_id = ? 
                ORDER BY r.submission_time DESC
            """, (telegram_id,))
            return [dict(row) for row in cursor.fetchall()]

    def has_participated(self, telegram_id, test_code):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT 1 FROM results WHERE telegram_id = ? AND test_code = ?", (telegram_id, test_code))
            return cursor.fetchone() is not None

    def get_test_results(self, test_code):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT r.*, u.full_name, u.phone, u.class_name 
                FROM results r 
                JOIN users u ON r.telegram_id = u.telegram_id 
                WHERE r.test_code = ? 
                ORDER BY r.percentage DESC, r.duration_spent ASC
            """, (test_code,))
            return [dict(row) for row in cursor.fetchall()]

    def get_user_rank(self, telegram_id, test_code):
        results = self.get_test_results(test_code)
        for rank, res in enumerate(results, 1):
            if res["telegram_id"] == telegram_id:
                return rank, len(results)
        return None, len(results)

    def get_test_leaderboard(self, test_code, limit=10):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT r.percentage, u.full_name, u.class_name
                FROM results r
                JOIN users u ON r.telegram_id = u.telegram_id
                WHERE r.test_code = ?
                ORDER BY r.percentage DESC, r.duration_spent ASC
                LIMIT ?
            """, (test_code, limit))
            return [dict(row) for row in cursor.fetchall()]

    # Statistics operations
    def get_statistics(self):
        stats = {}
        with self._get_connection() as conn:
            cursor = conn.cursor()
            
            # Total users
            cursor.execute("SELECT COUNT(*) FROM users")
            stats["total_users"] = cursor.fetchone()[0]
            
            # Active users (users who worked at least one test)
            cursor.execute("SELECT COUNT(DISTINCT telegram_id) FROM results")
            stats["active_users"] = cursor.fetchone()[0]
            
            # Total participants (total results)
            cursor.execute("SELECT COUNT(*) FROM results")
            stats["total_results"] = cursor.fetchone()[0]
            
            # Scores stats
            cursor.execute("SELECT MAX(percentage), MIN(percentage), AVG(percentage) FROM results")
            row = cursor.fetchone()
            if row and row[0] is not None:
                stats["highest_score"] = round(row[0], 1)
                stats["lowest_score"] = round(row[1], 1)
                stats["average_score"] = round(row[2], 1)
            else:
                stats["highest_score"] = 0
                stats["lowest_score"] = 0
                stats["average_score"] = 0
                
            # Certificates count
            cursor.execute("SELECT COUNT(*) FROM certificates")
            stats["certificates_count"] = cursor.fetchone()[0]
            
        return stats

    # Certificate operations
    def add_certificate(self, cert_id, telegram_id, test_code, pdf_path, image_path):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            issue_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor.execute("""
                INSERT OR REPLACE INTO certificates (id, telegram_id, test_code, issue_date, pdf_path, image_path)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (cert_id, telegram_id, test_code, issue_date, pdf_path, image_path))
            conn.commit()

    def get_certificate(self, cert_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM certificates WHERE id = ?", (cert_id,))
            row = cursor.fetchone()
            return dict(row) if row else None

    def get_user_certificates(self, telegram_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT c.*, t.subject, r.percentage, r.correct_count
                FROM certificates c
                JOIN tests t ON c.test_code = t.code
                JOIN results r ON c.telegram_id = r.telegram_id AND c.test_code = r.test_code
                WHERE c.telegram_id = ?
                ORDER BY c.issue_date DESC
            """, (telegram_id,))
            return [dict(row) for row in cursor.fetchall()]

    def get_global_leaderboard(self, limit=10):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT u.full_name, u.class_name, 
                       COUNT(r.id) as tests_count, 
                       ROUND(AVG(r.percentage), 1) as avg_score
                FROM users u
                JOIN results r ON u.telegram_id = r.telegram_id
                GROUP BY u.telegram_id
                ORDER BY avg_score DESC, tests_count DESC
                LIMIT ?
            """, (limit,))
            return [dict(row) for row in cursor.fetchall()]

    def get_user_global_rank(self, telegram_id):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT telegram_id, ROUND(AVG(percentage), 1) as avg_score, COUNT(id) as tests_count
                FROM results
                GROUP BY telegram_id
                ORDER BY avg_score DESC, tests_count DESC
            """)
            rows = cursor.fetchall()
            for rank, row in enumerate(rows, 1):
                if row["telegram_id"] == telegram_id:
                    return rank, len(rows)
            return None, len(rows)


==================================================
FILE: .\handlers\__init__.py
==================================================

from aiogram import Router
from .admin import router as admin_router
from .user import router as user_router

router = Router()
router.include_router(admin_router)
router.include_router(user_router)


==================================================
FILE: .\handlers\admin\export.py
==================================================

import logging
import os
from aiogram import Router, F
from aiogram.types import Message, CallbackQuery, FSInputFile
from aiogram.utils.keyboard import InlineKeyboardBuilder

from loader import bot, db
from config.config import Config
from utils.excel_export import export_test_results

router = Router()
logger = logging.getLogger("EduMasterBot.handlers.admin.export")

def is_admin_user(telegram_id: int) -> bool:
    return db.is_admin(telegram_id, Config.ADMIN_ID)

# --- 📥 NATIJALARNI YUKLAB OLISH (EXPORT) ---
@router.message(F.text == "📥 Natijalarni yuklab olish")
async def show_export_tests(message: Message):
    if not is_admin_user(message.from_user.id):
        return
        
    tests = db.get_all_tests()
    if not tests:
        await message.answer("❌ Bazada hech qanday test mavjud emas! Natijalarni eksport qilish uchun avval test qo'shing.")
        return
        
    kb = InlineKeyboardBuilder()
    for t in tests:
        # Show code and subject in button
        kb.button(text=f"{t['code']} - {t['subject']}", callback_data=f"exp_test:{t['code']}")
        
    kb.adjust(1)
    
    await message.answer(
        "📥 <b>Natijalarni yuklab olmoqchi bo'lgan testingizni tanlang:</b>",
        reply_markup=kb.as_markup()
    )

# Export test selected callback
@router.callback_query(F.data.startswith("exp_test:"))
async def process_export_test_select(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    test_code = callback.data.split(":")[1]
    test = db.get_test(test_code)
    
    if not test:
        await callback.answer("Test topilmadi!", show_alert=True)
        return
        
    # Ask format selection
    kb = InlineKeyboardBuilder()
    kb.button(text="📊 Excel (.xlsx) formatida", callback_data=f"exp_fmt:{test_code}:excel")
    kb.button(text="📝 CSV formatida", callback_data=f"exp_fmt:{test_code}:csv")
    kb.button(text="📄 PDF formatida", callback_data=f"exp_fmt:{test_code}:pdf")
    kb.adjust(1)
    
    await callback.message.edit_text(
        f"⚙️ <b>'{test_code}' ({test['subject']}) test natijalari uchun formatni tanlang:</b>",
        reply_markup=kb.as_markup()
    )
    await callback.answer()

# Format selected callback
@router.callback_query(F.data.startswith("exp_fmt:"))
async def process_export_format_select(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    parts = callback.data.split(":")
    test_code = parts[1]
    fmt = parts[2]
    
    test = db.get_test(test_code)
    if not test:
        await callback.answer("Test topilmadi!", show_alert=True)
        return
        
    # Get all results for this test
    results = db.get_test_results(test_code)
    if not results:
        await callback.answer("⚠️ Ushbu test bo'yicha hali hech kim ishtirok etmagan! Eksport qilib bo'lmaydi.", show_alert=True)
        return
        
    await callback.answer("Fayl tayyorlanmoqda, iltimos kuting...", show_alert=False)
    
    try:
        # Generate the exported file
        file_path = export_test_results(results, test_code, test["subject"], fmt)
        
        if os.path.exists(file_path):
            doc = FSInputFile(file_path)
            # Send file to administrator
            caption = f"📊 <b>'{test_code}' ({test['subject']}) test natijalari</b>\nFormat: <code>{fmt.upper()}</code>\nIshtirokchilar: {len(results)} ta o'quvchi."
            await callback.message.answer_document(doc, caption=caption)
            await callback.message.delete()
        else:
            await callback.message.answer("⚠️ Fayl yaratishda xatolik yuz berdi. Iltimos qayta urinib ko'ring.")
    except Exception as e:
        logger.error(f"Error exporting results for test {test_code} in format {fmt}: {e}")
        await callback.message.answer(f"⚠️ Eksport xatoligi: {e}")


==================================================
FILE: .\handlers\admin\panel.py
==================================================

import logging
import asyncio
import os
from aiogram import Router, F
from aiogram.filters import Command
from aiogram.types import Message, ReplyKeyboardRemove, InlineKeyboardMarkup, InlineKeyboardButton, CallbackQuery, FSInputFile
from aiogram.fsm.context import FSMContext
from aiogram.utils.keyboard import InlineKeyboardBuilder

from loader import bot, db
from config.config import Config
from states.states import AdminStates
from keyboards.reply.keyboards import get_admin_menu_keyboard, get_main_menu_keyboard, get_back_keyboard
from utils.excel_export import export_statistics_report, export_users_list


router = Router()
logger = logging.getLogger("EduMasterBot.handlers.admin.panel")

# Helper check for admin
def is_admin_user(telegram_id: int) -> bool:
    return db.is_admin(telegram_id, Config.ADMIN_ID)

@router.message(Command("admin"))
async def admin_panel_start(message: Message, state: FSMContext):
    await state.clear()
    telegram_id = message.from_user.id
    
    if not is_admin_user(telegram_id):
        await message.answer("⚠️ <b>Kechirasiz, siz admin emassiz!</b>\nBu bo'lim faqat tizim administratorlari uchun.")
        return
        
    await message.answer(
        "👨‍💻 <b>ADMINISTRATOR PANELIGA XUSH KELIBSIZ!</b>\n\n"
        "Quyidagi boshqaruv tugmalaridan foydalaning:",
        reply_markup=get_admin_menu_keyboard()
    )

# --- 📊 STATISTIKALAR ---
@router.message(F.text == "📊 Statistikalar")
async def show_admin_stats(message: Message):
    if not is_admin_user(message.from_user.id):
        return
        
    stats = db.get_statistics()
    
    stats_text = (
        "📊 <b>EDU MASTER BOT TIZIM STATISTIKASI</b>\n\n"
        f"👥 <b>Jami foydalanuvchilar:</b> {stats['total_users']} ta\n"
        f"🏃‍♂️ <b>Aktiv foydalanuvchilar:</b> {stats['active_users']} ta\n"
        f"📝 <b>Jami topshirilgan testlar:</b> {stats['total_results']} ta\n"
        f"📜 <b>Berilgan sertifikatlar:</b> {stats['certificates_count']} ta\n\n"
        f"🥇 <b>Eng yuqori ball:</b> {stats['highest_score']}%\n"
        f"🥈 <b>Eng past ball:</b> {stats['lowest_score']}%\n"
        f"📈 <b>O'rtacha o'zlashtirish:</b> {stats['average_score']}%"
    )
    
    kb = InlineKeyboardBuilder()
    kb.button(text="📥 Excel hisobot", callback_data="export_stats:excel")
    kb.button(text="📥 PDF hisobot", callback_data="export_stats:pdf")
    kb.adjust(2)
    
    await message.answer(stats_text, reply_markup=kb.as_markup())

# --- 👥 FOYDALANUVCHILAR ---
@router.message(F.text == "👥 Foydalanuvchilar")
async def list_bot_users(message: Message):
    if not is_admin_user(message.from_user.id):
        return
        
    users = db.get_all_users()
    if not users:
        await message.answer("👥 Foydalanuvchilar ro'yxati bo'sh!")
        return
        
    text = f"👥 <b>JAMI RO'YXATDAN O'TGAN FOYDALANUVCHILAR ({len(users)} ta):</b>\n\n"
    # Show the last 15 users to prevent telegram message length limits
    recent_users = users[-15:]
    for idx, u in enumerate(reversed(recent_users), 1):
        text += f"{idx}. <b>{u['full_name']}</b> | Sinf: {u['class_name']} | Tel: {u['phone']} | ID: <code>{u['telegram_id']}</code>\n"
        
    if len(users) > 15:
        text += f"\n<i>...va yana {len(users) - 15} ta foydalanuvchi bor. To'liq ro'yxatni yuklab olishingiz mumkin.</i>"
        
    kb = InlineKeyboardBuilder()
    kb.button(text="📥 Excel ro'yxat", callback_data="export_users:excel")
    kb.button(text="📥 PDF ro'yxat", callback_data="export_users:pdf")
    kb.adjust(2)
    
    await message.answer(text, reply_markup=kb.as_markup())

# --- EXPORT CALLBACKS ---
@router.callback_query(F.data.startswith("export_stats:"))
async def handle_export_stats(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    format_type = callback.data.split(":")[1]
    await callback.answer("⏳ Hisobot tayyorlanmoqda...", show_alert=False)
    
    try:
        stats = db.get_statistics()
        file_path = export_statistics_report(stats, format_type)
        if os.path.exists(file_path):
            doc = FSInputFile(file_path)
            await callback.message.reply_document(
                doc,
                caption=f"📊 Edu Master Bot tizim statistikasi ({format_type.upper()})"
            )
        else:
            await callback.message.answer("⚠️ Hisobot fayli topilmadi.")
    except Exception as e:
        logger.error(f"Error exporting stats: {e}")
        await callback.message.answer(f"⚠️ Xatolik yuz berdi: {e}")

@router.callback_query(F.data.startswith("export_users:"))
async def handle_export_users(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    format_type = callback.data.split(":")[1]
    await callback.answer("⏳ Ro'yxat tayyorlanmoqda...", show_alert=False)
    
    try:
        users = db.get_all_users()
        file_path = export_users_list(users, format_type)
        if os.path.exists(file_path):
            doc = FSInputFile(file_path)
            await callback.message.reply_document(
                doc,
                caption=f"👥 Edu Master Bot foydalanuvchilar ro'yxati ({format_type.upper()})"
            )
        else:
            await callback.message.answer("⚠️ Ro'yxat fayli topilmadi.")
    except Exception as e:
        logger.error(f"Error exporting users list: {e}")
        await callback.message.answer(f"⚠️ Xatolik yuz berdi: {e}")

# --- 📢 XABAR YUBORISH (BROADCAST) ---
@router.message(F.text == "📢 Xabar yuborish")
async def start_broadcast(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    await message.answer(
        "📢 <b>Barcha foydalanuvchilarga yuboriladigan xabar matnini yuboring:</b>\n\n"
        "Xabar rasm, video, hujjat ko'rinishida yoki oddiy matn bo'lishi mumkin.",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_broadcast_msg)

@router.message(AdminStates.waiting_for_broadcast_msg)
async def process_broadcast_message(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    if message.text in ["🏠 Bosh menuga qaytish", "↩️ Ortga qaytish"]:
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return

        
    users = db.get_all_users()
    if not users:
        await message.answer("⚠️ Botda hech qanday foydalanuvchi yo'q!", reply_markup=get_admin_menu_keyboard())
        await state.clear()
        return
        
    status_msg = await message.answer(f"⏳ <b>Xabar yuborish boshlandi. Jami qabul qiluvchilar: {len(users)} ta...</b>")
    await state.clear()
    
    success_count = 0
    fail_count = 0
    
    for u in users:
        try:
            # Send duplicate copy of message to users
            await message.copy_to(chat_id=u["telegram_id"])
            success_count += 1
            # Add anti-flood delay (approx 30 messages per second maximum, so 0.05s sleep is extremely safe)
            await asyncio.sleep(0.05)
        except Exception as e:
            logger.warning(f"Failed to send broadcast to user {u['telegram_id']}: {e}")
            fail_count += 1
            
    await status_msg.delete()
    await message.answer(
        "✅ <b>Xabar tarqatish muvaffaqiyatli yakunlandi!</b>\n\n"
        f"📥 Yetkazildi: <b>{success_count} ta</b> foydalanuvchiga\n"
        f"❌ Yetkazilmadi (bloklanganlar): <b>{fail_count} ta</b>",
        reply_markup=get_admin_menu_keyboard()
    )

# --- 🔐 ADMINLAR BOSHQARUVI (ADMIN MANAGEMENT) ---
@router.message(F.text == "🔐 Adminlar")
async def manage_admins_menu(message: Message):
    if not is_admin_user(message.from_user.id):
        return
        
    # Get current admins
    admins = db.get_all_admins()
    
    text = (
        "🔐 <b>ADMINISTRATORLARNI BOSHQARISH BO'LIMI</b>\n\n"
        f"👑 <b>Asosiy Yaratuvchi (Owner):</b> <code>{Config.ADMIN_ID}</code>\n\n"
        "👤 <b>Tizimdagi qo'shimcha adminlar ro'yxati:</b>\n"
    )
    
    if admins:
        for idx, adm in enumerate(admins, 1):
            text += f"▪️ {idx}. ID: <code>{adm['telegram_id']}</code> (Qo'shilgan: <i>{adm['added_date']}</i>)\n"
    else:
        text += "<i>Hozircha qo'shimcha adminlar yo'q.</i>\n"
        
    # Inline buttons
    kb = InlineKeyboardBuilder()
    kb.button(text="➕ Yangi admin qo'shish", callback_data="manage_admin_add")
    kb.button(text="❌ Adminni o'chirish", callback_data="manage_admin_del")
    kb.adjust(1)
    
    await message.answer(text, reply_markup=kb.as_markup())

# Callback: Admin add trigger
@router.callback_query(F.data == "manage_admin_add")
async def callback_admin_add_trigger(callback: CallbackQuery, state: FSMContext):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    await callback.message.delete()
    await callback.message.answer(
        "🔐 <b>Yangi administrator qo'shish</b>\n\n"
        "Iltimos, yangi adminning <b>Telegram ID</b> raqamini yuboring:\n"
        "<i>(Masalan: <code>12345678</code>)</i>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_new_admin_id)
    await callback.answer()

# Process new admin ID input
@router.message(AdminStates.waiting_for_new_admin_id)
async def process_new_admin_id(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text in ["🏠 Bosh menuga qaytish", "↩️ Ortga qaytish"]:
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    if not input_text.isdigit():
        await message.answer(
            "⚠️ <b>Xatolik!</b>\n\nTelegram ID faqat raqamlardan iborat bo'lishi kerak. Iltimos raqam ko'rinishida yuboring:",
            reply_markup=get_back_keyboard()
        )
        return
        
    new_admin_id = int(input_text)

    
    # Check if already admin
    if db.is_admin(new_admin_id, Config.ADMIN_ID):
        await message.answer("⚠️ Ushbu foydalanuvchi allaqachon administrator ro'yxatida bor! Boshqa ID yuboring:")
        return
        
    # Add to DB
    db.add_admin(new_admin_id)
    await state.clear()
    
    await message.answer(
        f"✅ <b>Telegram ID <code>{new_admin_id}</code> muvaffaqiyatli administrator qilib qo'shildi!</b>",
        reply_markup=get_admin_menu_keyboard()
    )

# Callback: Admin delete list trigger
@router.callback_query(F.data == "manage_admin_del")
async def callback_admin_del_trigger(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    admins = db.get_all_admins()
    if not admins:
        await callback.answer("⚠️ O'chirish uchun qo'shimcha administratorlar mavjud emas!", show_alert=True)
        return
        
    kb = InlineKeyboardBuilder()
    for adm in admins:
        # Check user details if exists
        user_info = db.get_user(adm["telegram_id"])
        name_str = f" - {user_info['full_name']}" if user_info else ""
        kb.button(
            text=f"❌ O'chirish (ID: {adm['telegram_id']}{name_str})", 
            callback_data=f"del_admin_id:{adm['telegram_id']}"
        )
    kb.button(text="↩️ Bekor qilish", callback_data="manage_admin_cancel")
    kb.adjust(1)
    
    await callback.message.edit_text(
        "❌ <b>O'chirmoqchi bo'lgan administratoringizni tanlang:</b>",
        reply_markup=kb.as_markup()
    )
    await callback.answer()

# Callback: Delete specific admin ID
@router.callback_query(F.data.startswith("del_admin_id:"))
async def callback_delete_admin_id(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return
        
    target_id = int(callback.data.split(":")[1])
    
    # Remove from database
    db.remove_admin(target_id)
    
    await callback.answer("✅ Administrator muvaffaqiyatli o'chirildi!", show_alert=True)
    await callback.message.delete()
    
    # Resend updated list
    # Emulate message wrapper to reuse manage_admins_menu
    class FakeMessage:
        def __init__(self, from_user, chat):
            self.from_user = from_user
            self.chat = chat
        async def answer(self, text, reply_markup=None):
            return await callback.message.answer(text, reply_markup=reply_markup)
            
    fake_msg = FakeMessage(callback.from_user, callback.message.chat)
    await manage_admins_menu(fake_msg)

# Callback: Cancel action
@router.callback_query(F.data == "manage_admin_cancel")
async def callback_manage_admin_cancel(callback: CallbackQuery):
    await callback.message.delete()
    
    # Resend updated list
    class FakeMessage:
        def __init__(self, from_user, chat):
            self.from_user = from_user
            self.chat = chat
        async def answer(self, text, reply_markup=None):
            return await callback.message.answer(text, reply_markup=reply_markup)
            
    fake_msg = FakeMessage(callback.from_user, callback.message.chat)
    await manage_admins_menu(fake_msg)
    await callback.answer()

# --- ⚙️ MAJBURIY KANALLAR ---
@router.message(F.text == "⚙️ Majburiy kanallar")
async def manage_channels(message: Message):
    if not is_admin_user(message.from_user.id):
        return
        
    channels = db.get_all_channels()
    
    text = "⚙️ <b>Majburiy a'zolik kanallarini boshqarish</b>\n\n"
    if channels:
        for idx, ch in enumerate(channels, 1):
            text += f"{idx}. <b>ID:</b> {ch['channel_id']} \n   <b>URL:</b> {ch['url']}\n"
    else:
        text += "<i>Hozircha hech qanday kanal qo'shilmagan.</i>\n"
        
    kb = InlineKeyboardBuilder()
    kb.button(text="➕ Kanal qo'shish", callback_data="add_channel")
    kb.button(text="❌ Kanal o'chirish", callback_data="del_channel_list")
    kb.adjust(2)
    
    await message.answer(text, reply_markup=kb.as_markup(), disable_web_page_preview=True)

@router.callback_query(F.data == "add_channel")
async def add_channel_start(callback: CallbackQuery, state: FSMContext):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return
    
    await callback.message.delete()
    await callback.message.answer(
        "➕ <b>Yangi kanal qo'shish</b>\n\n"
        "Kanalning username yoki ID sini kiriting (masalan: @edumaster_channel yoki -100123456789):",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_channel_id)

@router.message(AdminStates.waiting_for_channel_id)
async def add_channel_id(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    text = message.text.strip()
    if text in ["🏠 Bosh menuga qaytish", "↩️ Ortga qaytish"]:
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    await state.update_data(channel_id=text)
    await message.answer(
        "Kanalning taklif havolasini (URL) yuboring (masalan: https://t.me/edumaster_channel):",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_channel_url)

@router.message(AdminStates.waiting_for_channel_url)
async def add_channel_url(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    text = message.text.strip()
    if text in ["🏠 Bosh menuga qaytish", "↩️ Ortga qaytish"]:
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    data = await state.get_data()
    channel_id = data['channel_id']
    url = text
    
    # Professional check: verify if the bot is admin in the channel
    try:
        from loader import bot
        # Format channel_id properly
        check_id = channel_id
        if not check_id.startswith("@") and not check_id.startswith("-"):
            check_id = f"@{check_id}"
            
        bot_member = await bot.get_chat_member(chat_id=check_id, user_id=bot.id)
        if bot_member.status not in ["administrator", "creator"]:
            await message.answer(
                "⚠️ <b>Xatolik!</b>\n"
                "Bot ushbu kanalda administrator emas. Iltimos, avval botni kanalga administrator qilib qo'shing va qayta urinib ko'ring.",
                reply_markup=get_admin_menu_keyboard()
            )
            await state.clear()
            return
    except Exception as e:
        logger.error(f"Error verifying channel: {e}")
        await message.answer(
            "⚠️ <b>Xatolik!</b>\n"
            "Kanalni topib bo'lmadi yoki botning kanalga kirish huquqi yo'q. "
            "Kanal ID si (yoki username) to'g'riligini va bot kanalda admin ekanligini tekshiring.",
            reply_markup=get_admin_menu_keyboard()
        )
        await state.clear()
        return

    db.add_channel(channel_id, url)
    await state.clear()
    
    await message.answer(f"✅ Kanal muvaffaqiyatli qo'shildi va tekshirildi:\n<b>{channel_id}</b>", reply_markup=get_admin_menu_keyboard())

@router.callback_query(F.data == "del_channel_list")
async def del_channel_list(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat yo'q", show_alert=True)
        return
        
    channels = db.get_all_channels()
    if not channels:
        await callback.answer("Kanal yo'q!", show_alert=True)
        return
        
    kb = InlineKeyboardBuilder()
    for ch in channels:
        kb.button(text=f"❌ {ch['channel_id']}", callback_data=f"del_channel:{ch['channel_id']}")
    kb.button(text="↩️ Bekor qilish", callback_data="manage_admin_cancel")
    kb.adjust(1)
    
    await callback.message.edit_text("❌ <b>O'chirmoqchi bo'lgan kanalni tanlang:</b>", reply_markup=kb.as_markup())

@router.callback_query(F.data.startswith("del_channel:"))
async def del_channel_callback(callback: CallbackQuery):
    if not is_admin_user(callback.from_user.id):
        return
    channel_id = callback.data.split(":", 1)[1]
    db.remove_channel(channel_id)
    await callback.answer("Kanal o'chirildi", show_alert=True)
    await callback.message.delete()
    
    fake_msg = type("FakeMessage", (), {"from_user": callback.from_user, "answer": callback.message.answer})()
    await manage_channels(fake_msg)


==================================================
FILE: .\handlers\admin\test_manage.py
==================================================

import os
import logging
from datetime import datetime
from aiogram import Router, F
from aiogram.types import Message, ReplyKeyboardRemove
from aiogram.fsm.context import FSMContext

from loader import bot, db
from config.config import Config
from states.states import AdminStates
from keyboards.reply.keyboards import get_admin_menu_keyboard, get_back_keyboard
from handlers.user.test import parse_user_answers

router = Router()
logger = logging.getLogger("EduMasterBot.handlers.admin.test_manage")

def is_admin_user(telegram_id: int) -> bool:
    return db.is_admin(telegram_id, Config.ADMIN_ID)

# --- ➕ TEST CREATION FLOW ---
@router.message(F.text == "➕ Test qo'shish")
async def start_add_test(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    await state.update_data(is_gift=0)
    await message.answer(
        "📝 <b>Yangi test kodini kiriting:</b>\n"
        "<i>(Masalan: PM-21)</i>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_code)

@router.message(F.text == "🎁 Sovg'ali test qo'shish")
async def start_add_gift_test(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    await state.update_data(is_gift=1)
    await message.answer(
        "🎁 <b>Yangi SOVG'ALI test kodini kiriting:</b>\n"
        "<i>(Masalan: GIFT-01)</i>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_code)

@router.message(AdminStates.waiting_for_code)
async def process_code(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    code = input_text.upper()
    if not code.isalnum() and "-" not in code:
        await message.answer("⚠️ Test kodida faqat harflar, raqamlar va '-' belgisi qatnashishi mumkin. Qaytadan kiriting:", reply_markup=get_back_keyboard())
        return
        
    # Check if test code exists
    existing = db.get_test(code)
    if existing:
        await message.answer("⚠️ Ziddiyat! Ushbu kodga ega test allaqachon bazada bor. Iltimos boshqa kod yozing:", reply_markup=get_back_keyboard())
        return
        
    await state.update_data(code=code)
    await message.answer("📚 <b>Test topshiriladigan fanning nomini kiriting:</b>\n<i>(Masalan: Matematika)</i>", reply_markup=get_back_keyboard())
    await state.set_state(AdminStates.waiting_for_subject)

@router.message(AdminStates.waiting_for_subject)
async def process_subject(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    await state.update_data(subject=input_text)
    await message.answer("📄 <b>Test savollari yozilgan PDF faylni yuboring (yuklang):</b>", reply_markup=get_back_keyboard())
    await state.set_state(AdminStates.waiting_for_pdf)

@router.message(AdminStates.waiting_for_pdf, F.text == "↩️ Ortga qaytish")
async def cancel_pdf_upload(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())

@router.message(AdminStates.waiting_for_pdf, F.document)
async def process_pdf(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    doc = message.document
    if not doc.file_name.lower().endswith(".pdf"):
        await message.answer("⚠️ Xatolik! Iltimos, faqat PDF formatidagi fayl yuklang:", reply_markup=get_back_keyboard())
        return
        
    data = await state.get_data()
    test_code = data["code"]
    
    # Download file securely
    file_info = await bot.get_file(doc.file_id)
    pdf_filename = f"{test_code}.pdf"
    pdf_filepath = Config.PDF_DIR / pdf_filename
    
    Config.PDF_DIR.mkdir(parents=True, exist_ok=True)
    await bot.download_file(file_info.file_path, destination=pdf_filepath)
    
    await state.update_data(pdf_path=str(pdf_filepath))
    await message.answer("🎯 <b>To'g'ri javoblarni bir qator qilib yuboring:</b>\n<i>(Masalan: ABCDABCDABCD)</i>", reply_markup=get_back_keyboard())
    await state.set_state(AdminStates.waiting_for_answers)

@router.message(AdminStates.waiting_for_answers)
async def process_answers(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    parsed_answers = parse_user_answers(input_text)
    if not parsed_answers:
        await message.answer(
            "⚠️ <b>Javoblar aniqlanmadi!</b>\n\n"
            "Iltimos, to'g'ri javoblarni quyidagi to'g'ri formatlardan birida yuboring:\n"
            "👉 Ketma-ket: <code>ABCDABCD</code>\n"
            "👉 Vergul bilan: <code>1.a, 2.b, 3.c</code>\n"
            "👉 Tire bilan: <code>1-a 2-b 3-c</code>",
            reply_markup=get_back_keyboard()
        )
        return
        
    await state.update_data(answers=parsed_answers)

    await message.answer("⏳ <b>Test ishlash vaqtini daqiqalarda kiriting:</b>\n<i>(Masalan: 30)</i>", reply_markup=get_back_keyboard())
    await state.set_state(AdminStates.waiting_for_duration)

@router.message(AdminStates.waiting_for_duration)
async def process_duration(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    if not input_text.isdigit():
        await message.answer("⚠️ Vaqt faqat raqamda bo'lishi kerak. Masalan: 30 (daqiqa):", reply_markup=get_back_keyboard())
        return
        
    duration = int(input_text)
    await state.update_data(duration=duration)
    await message.answer(
        "📅 <b>Test boshlanish vaqtini kiriting:</b>\n"
        "Format: <code>DD.MM.YYYY HH:MM</code>\n"
        "<i>Masalan: 20.05.2026 10:00</i>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_start)

@router.message(AdminStates.waiting_for_start)
async def process_start_time(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    try:
        # Validate date format
        datetime.strptime(input_text, "%d.%m.%Y %H:%M")
    except ValueError:
        await message.answer(
            "⚠️ Vaqt formati noto'g'ri! Iltimos quyidagi formatda kiriting:\n"
            "<code>DD.MM.YYYY HH:MM</code> (Masalan: 20.05.2026 10:00):",
            reply_markup=get_back_keyboard()
        )
        return
        
    await state.update_data(start_time=input_text)
    await message.answer(
        "📅 <b>Test tugash vaqtini kiriting:</b>\n"
        "Format: <code>DD.MM.YYYY HH:MM</code>\n"
        "<i>Masalan: 20.05.2026 12:00</i>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(AdminStates.waiting_for_end)

@router.message(AdminStates.waiting_for_end)
async def process_end_time(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    try:
        # Validate date format
        datetime.strptime(input_text, "%d.%m.%Y %H:%M")
    except ValueError:
        await message.answer(
            "⚠️ Vaqt formati noto'g'ri! Iltimos quyidagi formatda kiriting:\n"
            "<code>DD.MM.YYYY HH:MM</code> (Masalan: 20.05.2026 12:00):",
            reply_markup=get_back_keyboard()
        )
        return
        
    data = await state.get_data()
    
    # Verify end time is after start time
    start_dt = datetime.strptime(data["start_time"], "%d.%m.%Y %H:%M")
    end_dt = datetime.strptime(input_text, "%d.%m.%Y %H:%M")
    
    if end_dt <= start_dt:
        await message.answer("⚠️ Xatolik! Test tugash vaqti boshlanish vaqtidan keyin bo'lishi kerak. Qaytadan kiriting:", reply_markup=get_back_keyboard())
        return
        
    # Save to SQLite
    is_gift = data.get("is_gift", 0)
    db.add_test(
        code=data["code"],
        subject=data["subject"],
        pdf_path=data["pdf_path"],
        correct_answers=data["answers"],
        duration=data["duration"],
        start_time=data["start_time"],
        end_time=input_text,
        is_gift=is_gift
    )
    
    await state.clear()
    
    gift_prefix = "🎁 SOVG'ALI " if is_gift else "✅ "
    confirm_msg = (
        f"{gift_prefix}<b>TEST MUVAFFAQIYATLI QO'SHILDI!</b>\n\n"
        f"📄 Test kodi: <b>{data['code']}</b>\n"
        f"📚 Fan: <b>{data['subject']}</b>\n"
        f"🎯 Javoblar soni: <b>{len(data['answers'])} ta</b>\n"
        f"⏳ Ishlash vaqti: <b>{data['duration']} daqiqa</b>\n"
        f"📅 Muddati: <b>{data['start_time']}</b> dan <b>{input_text}</b> gacha"
    )
    
    await message.answer(confirm_msg, reply_markup=get_admin_menu_keyboard())

    # Generate Promotional Advertisement Post with 4 text links
    promo_text = (
        f"📢 <b>YANGI TEST E'LON QILINDI!</b>\n\n"
        f"📚 <b>Fan:</b> {data['subject']}\n"
        f"🔑 <b>Test kodi:</b> <code>{data['code']}</code>\n"
        f"⏳ <b>Ishlash vaqti:</b> {data['duration']} daqiqa\n"
        f"📅 <b>Boshlanish:</b> {data['start_time']}\n"
        f"📅 <b>Tugash:</b> {input_text}\n\n"
        f"💡 <i>Ushbu testni ishlash uchun quyidagi bot havolalaridan biri orqali botga o'ting va test kodini kiriting:</i>\n\n"
        f"👉 @Edu_master2_bot\n"
        f"👉 @Edu_master2_bot\n"
        f"👉 @Edu_master2_bot\n"
        f"👉 @Edu_master2_bot"
    )
    
    await message.answer("👇 <b>Kanal yoki guruhlarga tarqatish (forward) uchun tayyor reklama posti:</b>")
    await message.answer(promo_text)

# --- ❌ TESTNI O'CHIRISH ---
@router.message(F.text == "❌ Testni o'chirish")
async def start_delete_test(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    tests = db.get_all_tests()
    if not tests:
        await message.answer("❌ Bazada hech qanday test mavjud emas!")
        return
        
    text = "📝 <b>Bazada mavjud testlar ro'yxati:</b>\n\n"
    for t in tests:
        text += f"▪️ Kodi: <code>{t['code']}</code> | Fan: {t['subject']}\n"
        
    text += "\n❌ <b>O'chirmoqchi bo'lgan test kodini yuboring:</b>"
    
    await message.answer(text, reply_markup=get_back_keyboard())
    await state.set_state(AdminStates.waiting_for_delete_code)

@router.message(AdminStates.waiting_for_delete_code)
async def process_delete_test(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    code = input_text.upper()
    test = db.get_test(code)
    
    if not test:
        await message.answer(
            "⚠️ Bunday kodli test topilmadi! Qaytadan kiriting:\n"
            "<i>(Ortga qaytish uchun: '↩️ Ortga qaytish' tugmasini bosing)</i>",
            reply_markup=get_back_keyboard()
        )
        return
        
    # Delete test file if exists
    if test["pdf_path"] and os.path.exists(test["pdf_path"]):
        try:
            os.remove(test["pdf_path"])
        except Exception as e:
            logger.error(f"Error removing PDF file: {e}")
            
    # Delete from DB
    db.delete_test(code)
    await state.clear()
    
    await message.answer(
        f"✅ <b>'{code}' kodli test va unga tegishli barcha natijalar, sertifikatlar tizimdan butunlay o'chirildi!</b>",
        reply_markup=get_admin_menu_keyboard()
    )

# --- ✏️ TESTNI TAHRIRLASH ---
@router.message(F.text == "✏️ Testni tahrirlash")
async def start_edit_test(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        return
        
    tests = db.get_all_tests()
    if not tests:
        await message.answer("❌ Bazada hech qanday test mavjud emas!")
        return
        
    text = "📝 <b>Bazada mavjud testlar ro'yxati:</b>\n\n"
    for t in tests:
        status = "Faol ✅" if t["is_active"] else "To'xtatilgan ❌"
        text += f"▪️ Kodi: <code>{t['code']}</code> | Fan: {t['subject']} | Holati: {status}\n"
        
    text += "\n✏️ <b>Tahrirlamoqchi bo'lgan test kodini yuboring:</b>"
    
    await message.answer(text, reply_markup=get_back_keyboard())
    await state.set_state(AdminStates.waiting_for_edit_code)

@router.message(AdminStates.waiting_for_edit_code)
async def process_edit_test_code(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    code = input_text.upper()
    test = db.get_test(code)
    
    if not test:
        await message.answer(
            "⚠️ Bunday kodli test topilmadi! Qaytadan kiriting:\n"
            "<i>(Ortga qaytish uchun: '↩️ Ortga qaytish' tugmasini bosing)</i>",
            reply_markup=get_back_keyboard()
        )
        return
        
    await state.update_data(edit_code=code)
    
    from aiogram.utils.keyboard import ReplyKeyboardBuilder
    kb = ReplyKeyboardBuilder()
    kb.button(text="⏱ Vaqtni tahrirlash")
    kb.button(text="🎯 Javoblarni tahrirlash")
    kb.button(text="🔄 Holatini o'zgartirish (Faol/To'xtatish)")
    kb.button(text="↩️ Ortga qaytish")
    kb.adjust(1)
    
    await message.answer(
        f"✏️ <b>{code}</b> kodli test uchun nimani tahrirlamoqchisiz?",
        reply_markup=kb.as_markup(resize_keyboard=True)
    )
    await state.set_state(AdminStates.waiting_for_edit_field)

@router.message(AdminStates.waiting_for_edit_field)
async def process_edit_field(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    if input_text == "⏱ Vaqtni tahrirlash":
        await state.update_data(edit_field="duration")
        await message.answer("⏳ <b>Testning yangi vaqtini daqiqalarda kiriting:</b>\n<i>(Masalan: 45)</i>", reply_markup=get_back_keyboard())
        await state.set_state(AdminStates.waiting_for_edit_value)
    elif input_text == "🎯 Javoblarni tahrirlash":
        await state.update_data(edit_field="answers")
        await message.answer("🎯 <b>Yangi to'g'ri javoblarni bir qator qilib yuboring:</b>\n<i>(Masalan: ABCDABCDABCD)</i>", reply_markup=get_back_keyboard())
        await state.set_state(AdminStates.waiting_for_edit_value)
    elif input_text == "🔄 Holatini o'zgartirish (Faol/To'xtatish)":
        data = await state.get_data()
        test = db.get_test(data["edit_code"])
        new_status = 0 if test["is_active"] else 1
        db.update_test_status(data["edit_code"], new_status)
        status_text = "Faol" if new_status else "To'xtatilgan"
        await state.clear()
        await message.answer(f"✅ Test holati <b>{status_text}</b> holatiga o'zgartirildi!", reply_markup=get_admin_menu_keyboard())
    else:
        await message.answer("⚠️ Noto'g'ri buyruq tanlandi. Klaviaturadagi tugmalardan foydalaning.")

@router.message(AdminStates.waiting_for_edit_value)
async def process_edit_value(message: Message, state: FSMContext):
    if not is_admin_user(message.from_user.id):
        await state.clear()
        return
        
    input_text = message.text.strip()
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_admin_menu_keyboard())
        return
        
    data = await state.get_data()
    code = data["edit_code"]
    field = data["edit_field"]
    
    if field == "duration":
        if not input_text.isdigit():
            await message.answer("⚠️ Vaqt faqat raqamda bo'lishi kerak. Masalan: 30 (daqiqa):")
            return
        db.update_test(code, duration=int(input_text))
        await message.answer(f"✅ Test ishlash vaqti <b>{input_text} daqiqa</b> qilib o'zgartirildi!", reply_markup=get_admin_menu_keyboard())
        
    elif field == "answers":
        parsed = parse_user_answers(input_text)
        if not parsed:
            await message.answer("⚠️ Javoblar aniqlanmadi! To'g'ri formatda yuboring:")
            return
        db.update_test(code, correct_answers=parsed)
        await message.answer(f"✅ Test javoblari muvaffaqiyatli yangilandi!\nJavoblar: <code>{parsed}</code>", reply_markup=get_admin_menu_keyboard())
        
    await state.clear()


==================================================
FILE: .\handlers\admin\__init__.py
==================================================

from aiogram import Router
from . import panel, test_manage, export

router = Router()
router.include_router(panel.router)
router.include_router(test_manage.router)
router.include_router(export.router)


==================================================
FILE: .\handlers\user\menu.py
==================================================

import os
import logging
from datetime import datetime

from aiogram import Router, F
from aiogram.types import Message, CallbackQuery, ReplyKeyboardRemove
from aiogram.fsm.context import FSMContext

from loader import bot, db
from config.config import Config
from states.states import ProfileStates, RatingStates

from keyboards.reply.keyboards import get_main_menu_keyboard, get_class_keyboard, get_contact_keyboard, get_back_keyboard, get_admin_menu_keyboard
from keyboards.inline.keyboards import get_profile_edit_keyboard
from utils.subscription import is_subscribed

router = Router()
logger = logging.getLogger("EduMasterBot.handlers.menu")

from aiogram.filters import StateFilter

@router.message(F.text.in_(["↩️ Ortga qaytish", "↩️ Ortga", "🏠 Bosh menuga qaytish"]), StateFilter("*"))
async def global_back_handler(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())

# Middleware check function to ensure subscription
async def check_user_access(message: Message) -> bool:
    telegram_id = message.from_user.id
    user = db.get_user(telegram_id)
    if not user:
        await message.answer("⚠️ Botdan foydalanish uchun iltimos /start buyrug'ini bosing va ro'yxatdan o'ting.")
        return False
        
    channels = db.get_all_channels()
    not_subscribed = []
    if channels:
        for ch in channels:
            if not await is_subscribed(bot, ch['channel_id'], telegram_id):
                not_subscribed.append(ch)

    if not_subscribed:
        from keyboards.inline.keyboards import get_subscription_keyboard
        await message.answer(
            f"🔒 <b>Botdan foydalanish uchun hamkor kanallarimizga a'zo bo'lishingiz majburiy:</b>\n",
            reply_markup=get_subscription_keyboard(not_subscribed)
        )
        return False
    return True

# --- 👤 PROFILIM SECTION ---
@router.message(F.text == "👤 Profilim")
async def show_profile(message: Message):
    if not await check_user_access(message):
        return
        
    telegram_id = message.from_user.id
    user = db.get_user(telegram_id)
    results = db.get_user_results(telegram_id)
    
    msg_text = (
        "👤 <b>PROFIL MA'LUMOTLARI</b>\n\n"
        f"📝 <b>Ism va Familiya:</b> {user['full_name']}\n"
        f"📞 <b>Telefon raqam:</b> {user['phone']}\n"
        f"🏫 <b>Sinf:</b> {user['class_name']}\n"
        f"🆔 <b>Telegram ID:</b> <code>{user['telegram_id']}</code>\n"
        f"📅 <b>Ro'yxatdan o'tgan sana:</b> {user['registration_date']}\n"
        f"📚 <b>Topshirilgan testlar:</b> {len(results)} ta"
    )
    
    await message.answer(msg_text, reply_markup=get_profile_edit_keyboard())

# Profile edit callback
@router.callback_query(F.data == "edit_profile")
async def edit_profile_callback(callback: CallbackQuery, state: FSMContext):
    await callback.message.delete()
    await callback.message.answer(
        "📝 <b>Yangi ism va familiyangizni kiriting:</b>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(ProfileStates.waiting_for_new_name)
    await callback.answer()

@router.message(ProfileStates.waiting_for_new_name)
async def process_new_name(message: Message, state: FSMContext):
    name = message.text.strip()
    
    if name == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())
        return
        
    if len(name.split()) < 2:
        await message.answer(
            "⚠️ Iltimos, ism va familiyani to'liq kiriting (masalan: Ali Hasanov):",
            reply_markup=get_back_keyboard()
        )
        return
        
    await state.update_data(new_name=name)
    await message.answer(

        "🏫 <b>Yangi sinfingizni tanlang:</b>",
        reply_markup=get_class_keyboard()
    )
    await state.set_state(ProfileStates.waiting_for_new_class)

@router.message(ProfileStates.waiting_for_new_class)
async def process_new_class(message: Message, state: FSMContext):
    class_name = message.text.strip()
    
    if class_name == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())
        return
        
    valid_classes = [f"{i}-sinf" for i in range(1, 12)]
    if class_name not in valid_classes:
        await message.answer("⚠️ Iltimos, klaviaturadagi sinflardan birini tanlang:", reply_markup=get_class_keyboard())
        return
        
    user_data = await state.get_data()
    db.update_user(message.from_user.id, full_name=user_data["new_name"], class_name=class_name)
    await state.clear()
    
    await message.answer(
        "✅ <b>Profil ma'lumotlaringiz muvaffaqiyatli yangilandi!</b>",
        reply_markup=get_main_menu_keyboard()
    )

# Phone update callback
@router.callback_query(F.data == "update_phone")
async def update_phone_callback(callback: CallbackQuery, state: FSMContext):
    await callback.message.answer(
        "📞 <b>Yangi telefon raqamingizni yuboring:</b>\n"
        "Quyidagi tugmani bosing yoki raqamni quyidagi formatda kiriting: <i>+998901234567</i>",
        reply_markup=get_contact_keyboard()
    )
    await state.set_state(ProfileStates.waiting_for_new_phone)
    await callback.answer()

@router.message(ProfileStates.waiting_for_new_phone)
async def process_new_phone(message: Message, state: FSMContext):
    input_text = message.text.strip() if message.text else ""
    
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())
        return
        
    phone = ""
    if message.contact:
        phone = message.contact.phone_number
        if not phone.startswith("+"):
            phone = f"+{phone}"
    elif message.text:
        text = input_text
        if not text.startswith("+") and not text.isdigit():
            await message.answer(
                "⚠️ Noto'g'ri format. Telefon raqamingizni kiriting (+998901234567):",
                reply_markup=get_contact_keyboard()
            )
            return
        phone = text
        
    db.update_user(message.from_user.id, phone=phone)
    await state.clear()
    
    await message.answer(
        "✅ <b>Telefon raqamingiz muvaffaqiyatli yangilandi!</b>",
        reply_markup=get_main_menu_keyboard()
    )


# --- 📊 NATIJAM SECTION ---
@router.message(F.text == "📊 Natijam")
async def show_my_results(message: Message):
    if not await check_user_access(message):
        return
        
    telegram_id = message.from_user.id
    results = db.get_user_results(telegram_id)
    results = db.get_user_results(telegram_id)
    
    if not results:
        await message.answer(
            "📊 <b>Siz hali hech qanday test topshirmagansiz!</b>\n\n"
            "Test topshirish uchun <b>📚 Test ishlash</b> bo'limidan foydalaning."
        )
        return
        
    percentages = [r["percentage"] for r in results]
    avg_percentage = sum(percentages) / len(percentages)
    best_result = max(percentages)
    last_result = results[0]
    
    msg_text = (
        "🎯 <b>TEST NATIJALARINGIZ TAHLILI</b>\n\n"
        f"📊 <b>Umumiy topshirilgan testlar:</b> {len(results)} ta\n"
        f"🥇 <b>Eng yuqori natija:</b> {best_result:.1f}%\n"
        f"📈 <b>O'rtacha ko'rsatkich:</b> {avg_percentage:.1f}%\n\n"
        "📝 <b>Oxirgi topshirilgan test:</b>\n"
        f"▪️ Fan: {last_result['subject']}\n"
        f"▪️ Test kodi: <code>{last_result['test_code']}</code>\n"
        f"▪️ To'g'ri javoblar: {last_result['correct_count']} ta\n"
        f"▪️ Natija: {last_result['percentage']}%\n"
        f"▪️ Vaqt: {last_result['submission_time']}"
    )
    
    await message.answer(msg_text)

# --- 🏆 REYTING SECTION ---
@router.message(F.text == "🏆 Umumiy reyting")
async def show_rating_prompt(message: Message, state: FSMContext):
    if not await check_user_access(message):
        return
        
    telegram_id = message.from_user.id
    
    # Get overall global leaderboard (Top 10)
    leaderboard = db.get_global_leaderboard(limit=10)
    
    msg_text = "🏆 <b>EDU MASTER BOT GLOBAL REYTINGI (TOP 10)</b>\n"
    msg_text += "<i>Hamma foydalanuvchilarning ishlagan testlari va o'rtacha ballariga ko'ra:</i>\n\n"
    
    medals = ["🥇", "🥈", "🥉", "4️⃣", "5️⃣", "6️⃣", "7️⃣", "8️⃣", "9️⃣", "🔟"]
    
    if not leaderboard:
        msg_text += "ℹ️ Hozircha reytingda hech kim yo'q.\n\n"
    else:
        for idx, r in enumerate(leaderboard):
            icon = medals[idx] if idx < len(medals) else f"{idx+1}."
            name = r["full_name"]
            sinf = r["class_name"]
            avg_score = r["avg_score"]
            count = r["tests_count"]
            msg_text += f"{icon} <b>{name}</b> ({sinf}) — <b>{avg_score}%</b> <i>({count} ta test topshirgan)</i>\n"
            
    # Show user's own global rank
    u_rank, total_u = db.get_user_global_rank(telegram_id)
    msg_text += "\n-----------------------------------------\n"
    if u_rank:
        msg_text += f"👤 <b>Sizning o‘rningiz: {u_rank}-o‘rin</b> (jami {total_u} ta o‘quvchidan)\n\n"
    else:
        msg_text += "👤 <b>Siz hali biror test topshirmadingiz.</b>\n\n"
        
    msg_text += (
        "🔍 <b>Muayyan test bo‘yicha natijalar reytingini ko‘rishni istaysizmi?</b>\n"
        "Unda test kodini kiriting (masalan: <code>PM-21</code>):\n"
        "<i>(Bekor qilish uchun: '↩️ Ortga qaytish' tugmasini bosing)</i>"
    )
    
    await message.answer(msg_text, reply_markup=get_back_keyboard())
    await state.set_state(RatingStates.waiting_for_test_code)

@router.message(RatingStates.waiting_for_test_code)
async def process_rating_test_code(message: Message, state: FSMContext):
    if not await check_user_access(message):
        await state.clear()
        return
        
    input_text = message.text.strip()
    
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())
        return
        
    test_code = input_text.upper()
    telegram_id = message.from_user.id
    
    # Check test existence
    test = db.get_test(test_code)
    if not test:
        await message.answer(
            "❌ <b>Bunday kodli test topilmadi!</b>\n\n"
            "Iltimos, test kodini to‘g‘ri kiritganingizga ishonch hosil qiling va qayta yuboring:\n"
            "<i>(Ortga qaytish uchun: '↩️ Ortga qaytish' tugmasini bosing)</i>",
            reply_markup=get_back_keyboard()
        )
        return

        
    # Get test results sorted by percentage desc, duration spent asc
    results = db.get_test_results(test_code)
    
    # Clear state since input is processed
    await state.clear()
    
    if not results:
        await message.answer(
            f"🏆 <b>‘{test['subject']}’ ({test_code}) bo‘yicha hali natijalar mavjud emas.</b>\n\n"
            "Ushbu test bo‘yicha birinchi bo‘lib ishtirok eting!",
            reply_markup=get_main_menu_keyboard()
        )
        return
        
    msg_text = f"🏆 <b>‘{test['subject']}’ BO‘YICHA REYTING (TOP 3)</b>\n"
    msg_text += f"📄 Test kodi: <code>{test_code}</code>\n\n"
    
    medals = ["🥇", "🥈", "🥉"]
    for idx, r in enumerate(results[:3], 1):
        icon = medals[idx-1]
        name = r["full_name"]
        sinf = r["class_name"]
        score = r["percentage"]
        
        row_text = f"{icon} {name} ({sinf}) — <b>{score}%</b>\n"
        msg_text += row_text
        
    msg_text += "\n-----------------------------------------\n"
    
    # Get current user rank in this test
    rank, total_p = db.get_user_rank(telegram_id, test_code)
    if rank is not None:
        user_result = None
        for r in results:
            if r["telegram_id"] == telegram_id:
                user_result = r
                break
        score_str = f" ({user_result['percentage']}%)" if user_result else ""
        msg_text += f"👤 <b>Sizning o‘rningiz: {rank}-o‘rin{score_str}</b> (jami {total_p} ta o‘quvchidan)\n"
    else:
        msg_text += "👤 <b>Siz ushbu testda hali ishtirok etmadingiz.</b>\n"
        
    await message.answer(msg_text, reply_markup=get_main_menu_keyboard())


# --- ℹ️ YORDAM SECTION ---
@router.message(F.text == "ℹ️ Yordam")
async def show_help(message: Message):
    help_text = (
        "✨ <b>EDU MASTER — PROFESSIONAL YORDAM TIZIMI</b> ✨\n\n"
        "🤖 <b>Botdan to‘g‘ri va samarali foydalanish yo‘riqnomasi:</b>\n\n"
        "1️⃣ <b>📚 Test ishlash bo‘limi:</b>\n"
        "Asosiy menyudagi <b>📚 Test ishlash</b> tugmasini bosing va o‘zingiz topshirmoqchi bo‘lgan maxsus test kodini (masalan: <code>PM-21</code>) yozib yuboring.\n\n"
        "2️⃣ <b>📄 Savollar va topshiriqlar:</b>\n"
        "Test faol bo‘lgan taqdirda, tizim sizga PDF formatdagi savollar faylini yuboradi. Vaqt darhol ketishni boshlaydi.\n\n"
        "3️⃣ <b>🚀 Testni yakunlash va javob topshirish:</b>\n"
        "Savollarni yechib bo‘lgandan so‘ng, pastdagi <b>'🏁 Ishlab bo'ldim'</b> tugmasini bosing va javoblaringizni quyidagi to'g'ri formatlardan birida yozib yuboring:\n"
        "👉 Ketma-ket: <code>ABCDABCD</code>\n"
        "👉 Vergul bilan: <code>1.a, 2.b, 3.c</code>\n"
        "👉 Tire bilan: <code>1-a 2-b 3-c</code>\n\n"
        "🏆 <b>Reyting tizimi:</b>\n"
        "O‘z bilimingizni sinab ko‘ring va do‘stlaringiz orasida nechanchi o‘rinda ekanligingizni <b>🏆 Umumiy reyting</b> bo‘limi orqali doimiy kuzatib boring.\n\n"
        "🤝 <b>Texnik qo‘llab-quvvatlash va takliflar:</b>\n"
        "Agar sizda biron bir savol, taklif yoki texnik muammolar yuzaga kelsa, loyiha rahbariga murojaat qilishingiz mumkin:\n"
        "📞 <b>Murojaat uchun administrator:</b> @xolikulovic7"
    )
    await message.answer(help_text)


# --- 🤝 HAMKORLIK SECTION ---
@router.message(F.text == "🤝 Hamkorlik")
async def show_partnership(message: Message):
    partnership_text = (
        "🤝 <b>HAMKORLIK VA REKLAMA BO'LIMI</b>\n\n"
        "📢 Hurmatli hamkorlar va tadbirkorlar!\n"
        "Agar siz <b>Edu Master Bot</b> auditoriyasiga o'zingizning reklama xabaringizni joylamoqchi bo'lsangiz yoki hamkorlik loyihalarini taklif qilmoqchi bo'lsangiz, biz hamkorlikka doim tayyormiz.\n\n"
        "📩 <b>Reklama va tijoriy hamkorlik bo'yicha bog'lanish:</b>\n"
        "👉 Administrator: @xolikulovic7\n\n"
        "<i>Iltimos, murojaatingizda taklifingiz haqida to'liq va aniq ma'lumot qoldiring. Tez orada siz bilan bog'lanamiz!</i>"
    )
    await message.answer(partnership_text)

# --- 📢 REKLAMA BERISH SECTION ---
@router.message(F.text == "📢 Reklama berish")
async def show_advertisement(message: Message):
    ad_text = (
        "📢 <b>REKLAMA BERISH</b>\n\n"
        "Siz o'z mahsulotlaringiz, xizmatlaringiz yoki kanallaringizni botimizdagi minglab o'quvchilarga reklama qilishingiz mumkin!\n\n"
        "Bizda quyidagi reklama turlari mavjud:\n"
        "1️⃣ <b>Majburiy a'zolik</b> — Botga kiruvchi har bir yangi foydalanuvchi sizning kanalingizga a'zo bo'ladi.\n"
        "2️⃣ <b>Xabar tarqatish (Broadcast)</b> — Barcha bot a'zolariga sizning reklama xabaringiz yuboriladi.\n"
        "3️⃣ <b>Homiylik (Sovg'ali testlar)</b> — Haftalik testlarga homiylik qilib, brendingizni tanitishingiz mumkin.\n\n"
        "💰 <b>Narxlar va shartlar bo'yicha ma'lumot olish uchun admin bilan bog'laning:</b>\n"
        "👉 Administrator: @xolikulovic7"
    )
    await message.answer(ad_text)

# Go back to main menu
@router.message(F.text == "🏠 Bosh menuga qaytish")
async def back_to_main_menu(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("🏠 <b>Bosh menyu:</b>", reply_markup=get_main_menu_keyboard())

# --- 🎁 SOVG'ALI TESTLAR SECTION ---
@router.message(F.text == "🎁 Sovg'ali testlar")
async def show_gift_tests(message: Message):
    if not await check_user_access(message):
        return
        
    gifts = db.get_gift_tests()
    
    if not gifts:
        await message.answer(
            "🎁 <b>Hozirda faol sovg'ali testlar mavjud emas!</b>\n\n"
            "<i>Har haftada bir marta yangi sovg'ali testlar qo'shiladi. Kanalimizni va botni muntazam kuzatib boring!</i>",
            reply_markup=get_main_menu_keyboard()
        )
        return
        
    msg_text = (
        "🎁 <b>HAFTALIK SOVG'ALI TESTLAR RO'YXATI</b>\n"
        "<i>Quyidagi faol sovg'ali testlarda ishtirok eting va o'z bilimingizni sinab ko'ring!</i>\n\n"
    )
    
    from aiogram.utils.keyboard import InlineKeyboardBuilder
    kb = InlineKeyboardBuilder()
    
    for g in gifts:
        msg_text += f"⭐ <b>Fan:</b> {g['subject']} | <b>Kod:</b> <code>{g['code']}</code>\n"
        msg_text += f"⏳ Ishlash vaqti: {g['duration']} daqiqa\n"
        msg_text += f"📅 Muddati: {g['start_time']} dan {g['end_time']} gacha\n\n"
        
        # Button to instantly start this specific gift test
        kb.button(text=f"🚀 {g['subject']} ({g['code']}) ni boshlash", callback_data=f"start_gift_test:{g['code']}")
        
    kb.adjust(1)
    
    await message.answer(msg_text, reply_markup=kb.as_markup())


==================================================
FILE: .\handlers\user\start.py
==================================================

import logging
from datetime import datetime
from aiogram import Router, F
from aiogram.filters import CommandStart
from aiogram.types import Message, CallbackQuery, ReplyKeyboardRemove
from aiogram.fsm.context import FSMContext

from loader import bot, db
from config.config import Config
from states.states import RegistrationStates
from keyboards.reply.keyboards import get_contact_keyboard, get_class_keyboard, get_main_menu_keyboard
from keyboards.inline.keyboards import get_subscription_keyboard
from utils.subscription import is_subscribed

router = Router()
logger = logging.getLogger("EduMasterBot.handlers.start")

@router.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    telegram_id = message.from_user.id
    
    # 1. Check mandatory channel subscription first
    channels = db.get_all_channels()
    not_subscribed = []
    if channels:
        for ch in channels:
            if not await is_subscribed(bot, ch['channel_id'], telegram_id):
                not_subscribed.append(ch)

    if not_subscribed:
        await message.answer(
            f"👋 <b>Edu Master Botga xush kelibsiz!</b>\n\n"
            f"🔒 Botdan foydalanish uchun hamkor kanallarimizga a'zo bo'lishingiz majburiy.\n\n"
            f"<i>A'zo bo'lgandan so'ng, pastdagi '✅ Tekshirish' tugmasini bosing!</i>",
            reply_markup=get_subscription_keyboard(not_subscribed)
        )
        return

    # 2. Check if user already exists
    user = db.get_user(telegram_id)
    if not user:
        await message.answer(
            "👋 <b>Edu Master Botga xush kelibsiz!</b>\n\n"
            "Tizimdan foydalanish uchun ro'yxatdan o'tishingiz kerak.\n\n"
            "📝 <b>Ism va familiyangizni kiriting:</b>\n"
            "<i>(Masalan: Bahodir Boliqulov)</i>",
            reply_markup=ReplyKeyboardRemove()
        )
        await state.set_state(RegistrationStates.waiting_for_name)
    else:
        await message.answer(
            f"👋 <b>Qaytganingizdan xursandmiz, {user['full_name']}!</b>\n\n"
            "📚 Kerakli bo'limni tanlang:",
            reply_markup=get_main_menu_keyboard()
        )

# Registration FSM: 1. Full name
@router.message(RegistrationStates.waiting_for_name)
async def process_name(message: Message, state: FSMContext):
    name = message.text.strip()
    if len(name.split()) < 2:
        await message.answer(
            "⚠️ <b>Xatolik!</b>\n\n"
            "Iltimos, ism va familiyangizni to'liq kiriting (kamida 2 ta so'z):\n"
            "<i>(Masalan: Bahodir Boliqulov)</i>"
        )
        return
        
    await state.update_data(full_name=name)
    await message.answer(
        "📞 <b>Telefon raqamingizni yuboring:</b>\n\n"
        "Quyidagi <b>'📞 Telefon raqamni yuborish'</b> tugmasini bosing:",
        reply_markup=get_contact_keyboard()
    )
    await state.set_state(RegistrationStates.waiting_for_phone)

# Registration FSM: 2. Phone number
@router.message(RegistrationStates.waiting_for_phone, F.contact)
@router.message(RegistrationStates.waiting_for_phone, F.text)
async def process_phone(message: Message, state: FSMContext):
    phone = ""
    if message.contact:
        phone = message.contact.phone_number
        if not phone.startswith("+"):
            phone = f"+{phone}"
    elif message.text:
        text = message.text.strip()
        # Basic validation
        if not text.startswith("+") and not text.isdigit():
            await message.answer(
                "⚠️ <b>Telefon raqami formati noto'g'ri!</b>\n\n"
                "Iltimos, pastdagi tugmani bosing yoki raqamni quyidagi formatda yozib yuboring:\n"
                "<i>+998901234567</i>"
            )
            return
        phone = text

    await state.update_data(phone=phone)
    await message.answer(
        "🏫 <b>Sinfingizni tanlang:</b>\n\n"
        "Quyidagi ro'yxatdan o'zingiz o'qiydigan sinfni tanlang:",
        reply_markup=get_class_keyboard()
    )
    await state.set_state(RegistrationStates.waiting_for_class)

# Registration FSM: 3. Class
@router.message(RegistrationStates.waiting_for_class, F.text)
async def process_class(message: Message, state: FSMContext):
    class_name = message.text.strip()
    # Simple check if class is one of the choices
    valid_classes = [f"{i}-sinf" for i in range(1, 12)]
    if class_name not in valid_classes:
        await message.answer(
            "⚠️ <b>Sinf noto'g'ri tanlandi!</b>\n\n"
            "Iltimos, klaviaturadagi sinflardan birini tanlang:",
            reply_markup=get_class_keyboard()
        )
        return
        
    user_data = await state.get_data()
    telegram_id = message.from_user.id
    
    # Save to Database
    db.add_user(
        telegram_id=telegram_id,
        full_name=user_data["full_name"],
        phone=user_data["phone"],
        class_name=class_name
    )
    
    await state.clear()
    
    await message.answer(
        "🎉 <b>Muvaffaqiyatli ro'yxatdan o'tdingiz!</b>",
        reply_markup=ReplyKeyboardRemove()
    )
    
    # Check subscription
    channels = db.get_all_channels()
    not_subscribed = []
    if channels:
        for ch in channels:
            if not await is_subscribed(bot, ch['channel_id'], telegram_id):
                not_subscribed.append(ch)

    if not not_subscribed:
        await message.answer(
            f"✨ <b>Xush kelibsiz, {user_data['full_name']}!</b>\n\n"
            "Botdan foydalanishga tayyorsiz. Quyidagi menyudan foydalaning:",
            reply_markup=get_main_menu_keyboard()
        )
    else:
        await message.answer(
            f"🔒 <b>Ro'yxatdan o'tish muvaffaqiyatli yakunlandi.</b>\n\n"
            f"Biroq, botdan foydalanishni boshlash uchun kanallarimizga a'zo bo'lishingiz shart:",
            reply_markup=get_subscription_keyboard(not_subscribed)
        )

# Callback handler for subscription checking
@router.callback_query(F.data == "check_subscription")
async def process_check_subscription(callback: CallbackQuery, state: FSMContext):
    telegram_id = callback.from_user.id
    
    channels = db.get_all_channels()
    not_subscribed = []
    if channels:
        for ch in channels:
            if not await is_subscribed(bot, ch['channel_id'], telegram_id):
                not_subscribed.append(ch)

    if not_subscribed:
        await callback.answer(
            "❌ Siz hali hamma kanallarga a'zo bo'lmadingiz! Iltimos, a'zo bo'lib qayta tekshirib ko'ring.",
            show_alert=True
        )
        return
        
    user = db.get_user(telegram_id)
    await callback.message.delete()
    if not user:
        # Start FSM registration
        await callback.message.answer(
            "✅ <b>Rahmat! A'zolik tasdiqlandi.</b>\n\n"
            "Tizimdan foydalanish uchun ro'yxatdan o'tishingiz kerak.\n\n"
            "📝 <b>Ism va familiyangizni kiriting:</b>\n"
            "<i>(Masalan: Bahodir Boliqulov)</i>",
            reply_markup=ReplyKeyboardRemove()
        )
        await state.set_state(RegistrationStates.waiting_for_name)
    else:
        await callback.message.answer(
            f"✅ <b>Rahmat! Kanaldan a'zolik tekshirildi.</b>\n\n"
            f"✨ Xush kelibsiz, {user['full_name']}!\n"
            f"Kerakli bo'limni tanlang:",
            reply_markup=get_main_menu_keyboard()
        )
    await callback.answer()


==================================================
FILE: .\handlers\user\test.py
==================================================

import os
import logging
import time
import asyncio

from datetime import datetime
from aiogram import Router, F
from aiogram.types import Message, CallbackQuery, FSInputFile
from aiogram.fsm.context import FSMContext

from loader import bot, db
from config.config import Config
from states.states import TestStates
from keyboards.reply.keyboards import get_main_menu_keyboard, get_back_keyboard, get_solving_keyboard
from keyboards.inline.keyboards import get_test_start_keyboard, get_result_keyboard
from utils.subscription import is_subscribed
from utils.certificate import create_certificate
import re

def parse_user_answers(input_text: str) -> str:
    if not input_text:
        return ""
    input_text = input_text.upper().strip()
    # Matches patterns like 1.A, 1-A, 1 A, 1A
    matches = re.findall(r"(\d+)[\s\.\-]*([A-Z])", input_text)
    if matches:
        # Sort by the question number to ensure exact sequence order
        sorted_matches = sorted(matches, key=lambda x: int(x[0]))
        return "".join(val for num, val in sorted_matches)
    else:
        # Fallback to extracting all alphabetic characters directly
        return "".join(c for c in input_text if c.isalpha())

router = Router()
logger = logging.getLogger("EduMasterBot.handlers.test")

# Active timer tasks to cancel them if user submits early
active_timers = {}

async def test_timer_countdown(user_id: int, test_code: str, duration_seconds: int, state: FSMContext):
    """
    Background timer task. Sleeps for the duration of the test, 
    and automatically stops it if not already completed.
    """
    try:
        await asyncio.sleep(duration_seconds)
        
        # Check if the user is still taking this test
        current_state = await state.get_state()
        if current_state in [TestStates.solving, TestStates.waiting_for_answers]:
            state_data = await state.get_data()
            if state_data.get("active_test_code") == test_code:
                # Retrieve test details
                test = db.get_test(test_code)
                correct_answers = test["correct_answers"]
                total_q = len(correct_answers)
                
                # Auto submit with 0 correct answers
                user = db.get_user(user_id)
                db.add_result(
                    telegram_id=user_id,
                    test_code=test_code,
                    correct_count=0,
                    wrong_count=total_q,
                    percentage=0.0,
                    duration_spent=duration_seconds
                )
                
                await state.clear()
                
                # Notify User
                await bot.send_message(
                    user_id,
                    "⌛ <b>Vaqt tugadi!</b>\n\n"
                    f"Siz '{test['subject']}' testini topshirishga berilgan vaqtdan kechikdingiz. "
                    "Javoblaringiz qabul qilinmadi, natijangiz: 0%.",
                    reply_markup=get_main_menu_keyboard()
                )
    except asyncio.CancelledError:
        pass
    except Exception as e:
        logger.error(f"Error in test timer for user {user_id}: {e}")

# Check subscription middleware helper
async def check_user_access(message: Message) -> bool:
    telegram_id = message.from_user.id
    user = db.get_user(telegram_id)
    if not user:
        await message.answer("⚠️ Tizimdan foydalanish uchun /start buyrug'ini bosing va ro'yxatdan o'ting.")
        return False
        
    subscribed = await is_subscribed(bot, Config.CHANNEL_USERNAME, telegram_id)
    if not subscribed:
        from keyboards.inline.keyboards import get_subscription_keyboard
        await message.answer(
            f"🔒 <b>Botdan foydalanish uchun hamkor kanalimizga a'zo bo'lishingiz majburiy:</b>\n"
            f"👉 {Config.CHANNEL_USERNAME}",
            reply_markup=get_subscription_keyboard(Config.CHANNEL_USERNAME)
        )
        return False
    return True

# --- 📚 TEST ISHLASH SECTION ---
@router.message(F.text == "📚 Test ishlash")
async def enter_test_menu(message: Message, state: FSMContext):
    if not await check_user_access(message):
        return
        
    await message.answer(
        "📚 <b>Test topshirish bo‘limi</b>\n\n"
        "Iltimos, topshirmoqchi bo‘lgan <b>test kodini</b> yuboring:\n"
        "<i>(Masalan: <code>PM-21</code>)</i>",
        reply_markup=get_back_keyboard()
    )
    await state.set_state(TestStates.waiting_for_code)


async def start_solving_test_for_user(user_id: int, test_code: str, state: FSMContext, message: Message):
    # 1. Fetch test
    test = db.get_test(test_code)
    if not test:
        await message.answer(
            "❌ <b>Bunday kodli test topilmadi!</b>\n\n"
            "Iltimos, test kodini to'g'ri kiritganingizga ishonch hosil qiling va qayta yuboring:\n"
            "<i>(Ortga qaytish uchun: '↩️ Ortga qaytish' tugmasini bosing)</i>",
            reply_markup=get_back_keyboard()
        )
        return
        
    # 2. Check if already participated
    if db.has_participated(user_id, test_code):
        await message.answer(
            "⚠️ <b>Ziddiyat!</b>\n\n"
            "Siz ushbu testni topshirib bo'lgansiz. Qoida tariqasida har bir testda faqat bir marta ishtirok etish mumkin."
        )
        await state.clear()
        return

    # 3. Check if test active
    if not test["is_active"]:
        await message.answer("⚠️ <b>Ushbu test hozirda faol emas!</b>")
        await state.clear()
        return

    # 4. Check dates/times
    try:
        now = datetime.now()
        start_dt = datetime.strptime(test["start_time"], "%d.%m.%Y %H:%M")
        end_dt = datetime.strptime(test["end_time"], "%d.%m.%Y %H:%M")
        
        if now < start_dt:
            await message.answer(
                "⏳ <b>Ushbu test hali boshlanmadi!</b>\n\n"
                f"📅 Boshlanish vaqti: <b>{test['start_time']}</b>\n"
                f"🕒 Hozirgi vaqt: <b>{now.strftime('%d.%m.%Y %H:%M')}</b>"
            )
            await state.clear()
            return
            
        if now > end_dt:
            await message.answer(
                "❌ <b>Ushbu test muddati tugagan!</b>\n\n"
                f"📅 Tugash vaqti: <b>{test['end_time']}</b>\n"
                f"🕒 Hozirgi vaqt: <b>{now.strftime('%d.%m.%Y %H:%M')}</b>"
            )
            await state.clear()
            return
    except Exception as e:
        logger.error(f"Error parsing test times: {e}")
        await message.answer("⚠️ Test vaqtlarini tekshirishda xatolik yuz berdi. Admin bilan bog'laning.")
        await state.clear()
        return

    # 5. Show test details and start test immediately
    gift_tag = "🎁 SOVG'ALI " if test.get("is_gift", 0) else ""
    test_info = (
        f"🚀 <b>{gift_tag}TEST BOSHLANDI!</b>\n\n"
        f"📚 Fan: <b>{test['subject']}</b>\n"
        f"📄 Test kodi: <b>{test_code}</b>\n"
        f"⏳ Ishlash vaqti: <b>{test['duration']} daqiqa</b>\n"
        f"📅 Boshlanish: {test['start_time']}\n"
        f"📅 Tugash: {test['end_time']}\n\n"
        "Savollarni yeching. Ishlab bo'lganingizdan so'ng pastdagi <b>'🏁 Ishlab bo'ldim'</b> tugmasini bosing va javoblaringizni yuboring!"
    )
    
    start_timestamp = time.time()
    duration_seconds = test["duration"] * 60
    
    await state.update_data(
        active_test_code=test_code,
        start_time=start_timestamp,
        duration_seconds=duration_seconds
    )
    await state.set_state(TestStates.solving)
    
    # Send PDF first
    if test["pdf_path"] and os.path.exists(test["pdf_path"]):
        pdf_file = FSInputFile(test["pdf_path"])
        await message.answer_document(pdf_file, caption=f"📄 {test['subject']} savollari.pdf")
        
    await message.answer(test_info, reply_markup=get_solving_keyboard())
    
    # Create background task for auto submission
    timer_task = asyncio.create_task(
        test_timer_countdown(user_id, test_code, duration_seconds, state)
    )
    active_timers[f"{user_id}_{test_code}"] = timer_task

@router.message(TestStates.waiting_for_code)
async def process_test_code(message: Message, state: FSMContext):
    input_text = message.text.strip()
    
    # Check for back action
    if input_text == "↩️ Ortga qaytish":
        await state.clear()
        await message.answer("🏠 Bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())
        return
        
    test_code = input_text.upper()
    telegram_id = message.from_user.id
    
    await start_solving_test_for_user(telegram_id, test_code, state, message)

# Handle gift test start from callback button
@router.callback_query(F.data.startswith("start_gift_test:"))
async def process_callback_gift_test_start(callback: CallbackQuery, state: FSMContext):
    test_code = callback.data.split(":")[1]
    telegram_id = callback.from_user.id
    
    await callback.message.delete()
    await start_solving_test_for_user(telegram_id, test_code, state, callback.message)
    await callback.answer()

# User sends random text while solving
@router.message(TestStates.solving, F.text != "🏁 Ishlab bo'ldim")
async def solving_reminder(message: Message):
    await message.answer(
        "⏳ <b>Siz hozir test yechish jarayonidasiz!</b>\n\n"
        "Savollarni yechib bo'lgach, iltimos pastdagi <b>'🏁 Ishlab bo'ldim'</b> tugmasini bosing.",
        reply_markup=get_solving_keyboard()
    )

# When they click finished solving
@router.message(TestStates.solving, F.text == "🏁 Ishlab bo'ldim")
async def user_finished_solving(message: Message, state: FSMContext):
    await state.set_state(TestStates.waiting_for_answers)
    await message.answer(
        "✍️ <b>Javoblaringizni quyidagi to'g'ri formatlardan birida yuboring:</b>\n\n"
        "👉 Ketma-ket: <code>ABCDABCD</code>\n"
        "👉 Vergul bilan: <code>1.a, 2.b, 3.c</code>\n"
        "👉 Tire bilan: <code>1-a 2-b 3-c</code>",
        reply_markup=get_back_keyboard()
    )

# Receive and process answers under waiting_for_answers
@router.message(TestStates.waiting_for_answers)
async def process_answers(message: Message, state: FSMContext):
    input_text = message.text.strip() if message.text else ""
    telegram_id = message.from_user.id
    
    state_data = await state.get_data()
    test_code = state_data["active_test_code"]
    start_time = state_data["start_time"]
    
    if input_text == "↩️ Ortga qaytish":
        # Cancel active timer
        timer_key = f"{telegram_id}_{test_code}"
        if timer_key in active_timers:
            active_timers[timer_key].cancel()
            del active_timers[timer_key]
        await state.clear()
        await message.answer("🏠 Test bekor qilindi va bosh menyuga qaytildi.", reply_markup=get_main_menu_keyboard())
        return
        
    answers = parse_user_answers(input_text)
    test = db.get_test(test_code)
    correct_answers = test["correct_answers"]
    total_q = len(correct_answers)
    
    # Validate input length and chars
    if not answers:
        await message.answer(
            "⚠️ <b>Javoblaringiz aniqlanmadi!</b>\n\n"
            "Iltimos, quyidagi to'g'ri formatlardan birida yuboring:\n"
            "👉 Ketma-ket: <code>ABCDABCD</code>\n"
            "👉 Vergul bilan: <code>1.a, 2.b, 3.c</code>\n"
            "👉 Tire bilan: <code>1-a 2-b 3-c</code>"
        )
        return
        
    if len(answers) != total_q:
        await message.answer(
            f"⚠️ <b>Javoblar soni mos kelmadi!</b>\n\n"
            f"Ushbu testda jami <b>{total_q} ta</b> savol bor. Siz esa <b>{len(answers)} ta</b> javob yubordingiz.\n"
            f"Iltimos, qaytadan tekshirib, to'liq yuboring:"
        )
        return
        
    # Cancel active timer now that validation is passed
    timer_key = f"{telegram_id}_{test_code}"
    if timer_key in active_timers:
        active_timers[timer_key].cancel()
        del active_timers[timer_key]
        
    duration_spent = int(time.time() - start_time)
    
    # Check answers
    correct_count = 0
    wrong_count = 0
    wrong_details = []
    
    for idx in range(total_q):
        user_ans = answers[idx]
        correct_ans = correct_answers[idx]
        if user_ans == correct_ans:
            correct_count += 1
        else:
            wrong_count += 1
            wrong_details.append(
                f"❓ <b>{idx+1}-savol:</b> Siz yuborgan javob: '<code>{user_ans}</code>', to'g'ri javob: '<code>{correct_ans}</code>'.\n"
                f"⚠️ <i>(Ushbu savolni yaxshilab ko'rib chiqing!)</i>"
            )
            
    percentage = round((correct_count / total_q) * 100, 1)
    
    # Save results
    db.add_result(
        telegram_id=telegram_id,
        test_code=test_code,
        correct_count=correct_count,
        wrong_count=wrong_count,
        percentage=percentage,
        duration_spent=duration_spent
    )
    
    user = db.get_user(telegram_id)
    rank, total_p = db.get_user_rank(telegram_id, test_code)
    
    # Clear FSM State
    await state.clear()
    
    # Result Summary Page
    result_text = (
        "🎯 <b>TEST NATIJASI</b>\n\n"
        f"👤 <b>Foydalanuvchi:</b> {user['full_name']}\n"
        f"🏫 <b>Sinf:</b> {user['class_name']}\n"
        f"📚 <b>Fan:</b> {test['subject']}\n"
        f"📄 <b>Test kodi:</b> <code>{test_code}</code>\n\n"
        f"✅ <b>To‘g‘ri javoblar:</b> {correct_count} ta\n"
        f"❌ <b>Noto‘g‘ri javoblar:</b> {wrong_count} ta\n"
        f"📊 <b>Natija:</b> {percentage}%\n"
        f"🏆 <b>Reyting:</b> {rank}-o‘rin (jami {total_p} ta o'quvchidan)\n\n"
        f"⏱ <b>Sarflangan vaqt:</b> {duration_spent // 60} daqiqa {duration_spent % 60} soniya"
    )
    
    await message.answer(
        result_text, 
        reply_markup=get_result_keyboard(test_code)
    )
    
    # Send wrong answers feedback
    if wrong_details:
        feedback_text = "❌ <b>XATO ISHLANGAN SAVOLLAR TAHLILI:</b>\n"
        feedback_text += "<i>Xatolaringiz ustida ishlab, ushbu savollarni yaxshilab ko'rib chiqing:</i>\n\n"
        feedback_text += "\n\n".join(wrong_details)
        await message.answer(feedback_text)




@router.callback_query(F.data.startswith("view_rating:"))
async def view_test_leaderboard(callback: CallbackQuery):
    test_code = callback.data.split(":")[1]
    telegram_id = callback.from_user.id
    
    test = db.get_test(test_code)
    results = db.get_test_results(test_code)
    
    if not results:
        await callback.answer("⚠️ Bu test bo'yicha ma'lumot topilmadi!", show_alert=True)
        return
        
    msg_text = f"🏆 <b>'{test['subject']}' BO'YICHA REYTING (TOP 10)</b>\n"
    msg_text += f"📄 Test kodi: <code>{test_code}</code>\n\n"
    
    medals = ["🥇", "🥈", "🥉"]
    for idx, r in enumerate(results[:10], 1):
        icon = medals[idx-1] if idx <= 3 else f"<b>{idx}.</b>"
        name = r["full_name"]
        sinf = r["class_name"]
        score = r["percentage"]
        
        row_text = f"{icon} {name} ({sinf}) — <b>{score}%</b>\n"
        if r["telegram_id"] == telegram_id:
            row_text = f"👤 <b>{row_text.strip()} (Siz)</b>\n"
        msg_text += row_text
        
    await callback.message.answer(msg_text)
    await callback.answer()

@router.callback_query(F.data == "new_test")
async def start_new_test(callback: CallbackQuery, state: FSMContext):
    await callback.message.delete()
    await enter_test_menu(callback.message, state)
    await callback.answer()


==================================================
FILE: .\handlers\user\__init__.py
==================================================

from aiogram import Router
from . import start, menu, test

router = Router()
router.include_router(start.router)
router.include_router(menu.router)
router.include_router(test.router)


==================================================
FILE: .\keyboards\inline\keyboards.py
==================================================

from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton

def get_subscription_keyboard(channels: list):
    inline_keyboard = []
    for idx, channel in enumerate(channels, 1):
        inline_keyboard.append([InlineKeyboardButton(text=f"📢 {idx}-kanalga o'tish", url=channel['url'])])
    
    inline_keyboard.append([InlineKeyboardButton(text="✅ Tekshirish", callback_data="check_subscription")])
    return InlineKeyboardMarkup(inline_keyboard=inline_keyboard)

def get_test_start_keyboard(test_code: str):
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="🚀 Testni boshlash", callback_data=f"start_test:{test_code}")]
        ]
    )

def get_result_keyboard(test_code: str):
    buttons = []
    buttons.append([InlineKeyboardButton(text="🏆 Reytingni ko‘rish", callback_data=f"view_rating:{test_code}")])
    buttons.append([InlineKeyboardButton(text="📚 Yangi test", callback_data="new_test")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

def get_profile_edit_keyboard():
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✏️ Profilni tahrirlash", callback_data="edit_profile"),
                InlineKeyboardButton(text="📞 Telefonni yangilash", callback_data="update_phone")
            ]
        ]
    )

def get_admin_test_keyboard(test_code: str):
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="❌ O'chirish", callback_data=f"del_test:{test_code}")]
        ]
    )


==================================================
FILE: .\keyboards\reply\keyboards.py
==================================================

from aiogram.types import ReplyKeyboardMarkup, KeyboardButton

def get_contact_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="📞 Telefon raqamni yuborish", request_contact=True)],
            [KeyboardButton(text="↩️ Ortga qaytish")]
        ],
        resize_keyboard=True,
        one_time_keyboard=True
    )

def get_back_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="↩️ Ortga qaytish")]
        ],
        resize_keyboard=True
    )

def get_class_keyboard():
    # Buttons for 1-sinf to 11-sinf
    buttons = []
    row = []
    for i in range(1, 12):
        row.append(KeyboardButton(text=f"{i}-sinf"))
        if len(row) == 3:
            buttons.append(row)
            row = []
    if row:
        buttons.append(row)
    # Add Back button at the bottom
    buttons.append([KeyboardButton(text="↩️ Ortga qaytish")])
    return ReplyKeyboardMarkup(keyboard=buttons, resize_keyboard=True, one_time_keyboard=True)


def get_main_menu_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="📚 Test ishlash"), KeyboardButton(text="🎁 Sovg'ali testlar")],
            [KeyboardButton(text="📊 Natijam"), KeyboardButton(text="🏆 Umumiy reyting"), KeyboardButton(text="👤 Profilim")],
            [KeyboardButton(text="📢 Reklama berish"), KeyboardButton(text="ℹ️ Yordam"), KeyboardButton(text="🤝 Hamkorlik")]
        ],
        resize_keyboard=True
    )

def get_solving_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="🏁 Ishlab bo'ldim")]
        ],
        resize_keyboard=True
    )


def get_admin_menu_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="➕ Test qo'shish"), KeyboardButton(text="🎁 Sovg'ali test qo'shish")],
            [KeyboardButton(text="❌ Testni o'chirish"), KeyboardButton(text="✏️ Testni tahrirlash")],
            [KeyboardButton(text="📊 Statistikalar"), KeyboardButton(text="📥 Natijalarni yuklab olish")],
            [KeyboardButton(text="👥 Foydalanuvchilar"), KeyboardButton(text="📢 Xabar yuborish")],
            [KeyboardButton(text="⚙️ Majburiy kanallar"), KeyboardButton(text="🔐 Adminlar")],
            [KeyboardButton(text="🏠 Bosh menuga qaytish")]
        ],
        resize_keyboard=True
    )


==================================================
FILE: .\states\states.py
==================================================

from aiogram.fsm.state import State, StatesGroup

class RegistrationStates(StatesGroup):
    waiting_for_name = State()
    waiting_for_phone = State()
    waiting_for_class = State()

class ProfileStates(StatesGroup):
    waiting_for_new_name = State()
    waiting_for_new_class = State()
    waiting_for_new_phone = State()

class TestStates(StatesGroup):
    waiting_for_code = State()
    solving = State()  # Timer running, solving questions
    waiting_for_answers = State()  # User clicked "Ishlab bo'ldim" and is now sending their answers

class AdminStates(StatesGroup):
    # Test creation flow
    waiting_for_code = State()
    waiting_for_subject = State()
    waiting_for_pdf = State()
    waiting_for_answers = State()
    waiting_for_duration = State()
    waiting_for_start = State()
    waiting_for_end = State()
    
    # Broadcast message
    waiting_for_broadcast_msg = State()
    
    # Delete test
    waiting_for_delete_code = State()
    
    # Admin management
    waiting_for_new_admin_id = State()
    waiting_for_del_admin_id = State()
    
    # Test editing flow
    waiting_for_edit_code = State()
    waiting_for_edit_field = State()
    waiting_for_edit_value = State()
    
    # Channel management
    waiting_for_channel_id = State()
    waiting_for_channel_url = State()
    waiting_for_del_channel = State()

class RatingStates(StatesGroup):
    waiting_for_test_code = State()




==================================================
FILE: .\utils\certificate.py
==================================================

import os
import math
import random
import string
import qrcode
from PIL import Image, ImageDraw, ImageFont
from datetime import datetime
from pathlib import Path
from config.config import Config

def generate_certificate_id() -> str:
    # Generate a unique certificate ID, e.g. EMB-2026-00521
    year = datetime.now().year
    num = "".join(random.choices(string.digits, k=5))
    return f"EMB-{year}-{num}"

def create_certificate(full_name: str, class_name: str, subject: str, test_code: str, correct: int, total: int, percentage: float) -> dict:
    """
    Generates a premium education-style certificate matching the requested template exactly.
    Returns a dictionary with 'cert_id', 'image_path', and 'pdf_path'.
    """
    cert_id = generate_certificate_id()
    
    # 1. Create a base high-definition canvas (1280 x 960) with premium light cream linen background
    width, height = 1280, 960
    background_color = "#FDFBF7"  # Elegant off-white/cream
    image = Image.new("RGB", (width, height), background_color)
    draw = ImageDraw.Draw(image)
    
    # Color Palette
    navy_blue = "#083262"  # Premium dark navy blue
    gold = "#D2AC67"       # Sleek gold color
    dark_gray = "#333333"  # Text color
    
    # 2. Draw Elegant Repeating star watermarks for premium certificate paper look
    for wx in range(100, width, 180):
        for wy in range(100, height, 180):
            draw.polygon([
                (wx, wy - 6), (wx + 2, wy - 2), (wx + 6, wy - 2), (wx + 3, wy),
                (wx + 4, wy + 4), (wx, wy + 2), (wx - 4, wy + 4), (wx - 3, wy),
                (wx - 6, wy - 2), (wx - 2, wy - 2)
            ], fill="#FAF5EB")
            
    # Draw Subtle Wavy Abstract Silk Curves Across Canvas
    for offset in range(0, 200, 20):
        draw.arc([30 - offset, height - 450 - offset, 650 - offset, height - 30 - offset], 90, 180, fill="#F2E6D5", width=1)
        draw.arc([width - 650 + offset, 30 + offset, width - 30 + offset, 450 + offset], 270, 360, fill="#F2E6D5", width=1)
    
    # 3. Draw Premium Geometric Polygonal Overlays in corners with Gold Borders
    # Top-Right Corner
    draw.polygon([(width - 380, 40), (width - 40, 40), (width - 40, 380)], fill=navy_blue)
    draw.polygon([(width - 430, 40), (width - 380, 40), (width - 40, 380), (width - 40, 430)], fill=gold)
    
    # Bottom-Left Corner
    draw.polygon([(40, height - 380), (40, height - 40), (380, height - 40)], fill=navy_blue)
    draw.polygon([(40, height - 430), (40, height - 380), (380, height - 40), (430, height - 40)], fill=gold)
    
    # Bottom-Right Corner
    draw.polygon([(width - 480, height - 40), (width - 40, height - 480), (width - 40, height - 40)], fill=navy_blue)
    draw.polygon([(width - 530, height - 40), (width - 480, height - 40), (width - 40, height - 480), (width - 40, height - 530)], fill=gold)

    # 4. Draw Premium Double Gold and Navy Inner Borders
    # Primary gold outer border
    draw.rectangle([40, 40, width - 40, height - 40], outline=gold, width=3)
    # Secondary thin gold inner rim
    draw.rectangle([46, 46, width - 46, height - 46], outline=gold, width=1)
    # Inner navy blue frame
    draw.rectangle([50, 50, width - 50, height - 50], outline=navy_blue, width=2)
    
    # 5. Load System Fonts
    try:
        title_font = ImageFont.truetype("georgiab.ttf", 66)       # Georgia Bold
        subtitle_font = ImageFont.truetype("georgiai.ttf", 22)    # Georgia Italic
        cursive_font = ImageFont.truetype("gabriola.ttf", 75)      # Gabriola Cursive
        bold_serif = ImageFont.truetype("georgiab.ttf", 26)       # Georgia Bold small
        sans_bold = ImageFont.truetype("arialbd.ttf", 20)         # Arial Bold
        sans_regular = ImageFont.truetype("arial.ttf", 18)        # Arial Regular
        caption_font = ImageFont.truetype("arial.ttf", 14)        # Arial Caption
    except IOError:
        # Fallbacks
        title_font = ImageFont.load_default(50)
        subtitle_font = ImageFont.load_default(18)
        cursive_font = ImageFont.load_default(36)
        bold_serif = ImageFont.load_default(20)
        sans_bold = ImageFont.load_default(16)
        sans_regular = ImageFont.load_default(14)
        caption_font = ImageFont.load_default(12)
        
    # 6. Top-Left Premium Ribbed Golden Seal (Top-Left of Certificate)
    seal_cx, seal_cy = 160, 220
    
    # Draw ribbons hanging from seal
    draw.polygon([(seal_cx - 30, seal_cy + 10), (seal_cx - 50, seal_cy + 140), (seal_cx - 20, seal_cy + 115), (seal_cx - 5, seal_cy + 10)], fill=navy_blue)
    draw.polygon([(seal_cx - 30, seal_cy + 10), (seal_cx - 50, seal_cy + 140), (seal_cx - 20, seal_cy + 115), (seal_cx - 5, seal_cy + 10)], outline=gold, width=2)
    
    draw.polygon([(seal_cx + 5, seal_cy + 10), (seal_cx + 20, seal_cy + 115), (seal_cx + 50, seal_cy + 140), (seal_cx + 30, seal_cy + 10)], fill=navy_blue)
    draw.polygon([(seal_cx + 5, seal_cy + 10), (seal_cx + 20, seal_cy + 115), (seal_cx + 50, seal_cy + 140), (seal_cx + 30, seal_cy + 10)], outline=gold, width=2)
    
    # Generate jagged golden teeth for the outer seal ring
    teeth_points = []
    num_teeth = 40
    for i in range(2 * num_teeth):
        angle = i * math.pi / num_teeth
        r = 85 if i % 2 == 0 else 73
        x = seal_cx + r * math.cos(angle)
        y = seal_cy + r * math.sin(angle)
        teeth_points.append((x, y))
    draw.polygon(teeth_points, fill=gold)
    
    # Inner navy blue ring
    draw.ellipse([seal_cx - 62, seal_cy - 62, seal_cx + 62, seal_cy + 62], fill=navy_blue)
    # Inner gold border ring
    draw.ellipse([seal_cx - 55, seal_cy - 55, seal_cx + 55, seal_cy + 55], outline=gold, width=2)
    
    # Draw perfect programmatic gold star in seal center (instead of 🏆 emoji)
    star_pts = []
    for i in range(10):
        angle = i * math.pi / 5 - math.pi / 2
        r = 30 if i % 2 == 0 else 14
        star_pts.append((seal_cx + r * math.cos(angle), seal_cy + r * math.sin(angle)))
    draw.polygon(star_pts, fill=gold)
    
    # 7. Draw Top-Center Header with Programmatic graduation cap shape (instead of 🎓 emoji)
    cap_cx, cap_cy = width // 2, 105
    # Diamond top
    draw.polygon([(cap_cx, cap_cy - 12), (cap_cx + 24, cap_cy), (cap_cx, cap_cy + 12), (cap_cx - 24, cap_cy)], fill=navy_blue)
    draw.polygon([(cap_cx, cap_cy - 12), (cap_cx + 24, cap_cy), (cap_cx, cap_cy + 12), (cap_cx - 24, cap_cy)], outline=gold, width=2)
    # Base skull arch
    draw.chord([cap_cx - 12, cap_cy, cap_cx + 12, cap_cy + 16], 0, 180, fill=navy_blue)
    # Tassel
    draw.line([(cap_cx, cap_cy), (cap_cx - 16, cap_cy + 6)], fill=gold, width=2)
    draw.ellipse([cap_cx - 18, cap_cy + 5, cap_cx - 14, cap_cy + 9], fill=gold)
    
    draw.text((width // 2, 160), "EDU MASTER BOT", font=bold_serif, fill=navy_blue, anchor="mm")
    draw.text((width // 2, 190), "BILIM – KELAJAK KALITI!", font=caption_font, fill=gold, anchor="mm")
    
    # 8. Draw "S E R T I F I K A T" Title
    draw.text((width // 2, 280), "S E R T I F I K A T", font=title_font, fill=navy_blue, anchor="mm")
    
    # 9. Draw "Ushbu sertifikat" Gold Banner Ribbon
    banner_left, banner_top = width // 2 - 180, 335
    banner_right, banner_bottom = width // 2 + 180, 375
    draw.rectangle([banner_left, banner_top, banner_right, banner_bottom], fill=gold)
    draw.polygon([(banner_left, banner_top), (banner_left - 15, banner_top + 20), (banner_left, banner_bottom)], fill=gold)
    draw.polygon([(banner_right, banner_top), (banner_right + 15, banner_top + 20), (banner_right, banner_bottom)], fill=gold)
    draw.text((width // 2, 355), "Ushbu sertifikat", font=sans_bold, fill="#FFFFFF", anchor="mm")
    
    # 10. Draw Student Full Name in beautiful calligraphy style
    draw.text((width // 2, 445), full_name, font=cursive_font, fill=navy_blue, anchor="mm")
    
    # Elegant custom gold flourishes flanking the student name
    # Left flourish
    draw.arc([width // 2 - 290, 428, width // 2 - 240, 458], 0, 360, fill=gold, width=2)
    draw.line([(width // 2 - 250, 443), (width // 2 - 180, 443)], fill=gold, width=2)
    # Right flourish
    draw.arc([width // 2 + 240, 428, width // 2 + 290, 458], 0, 360, fill=gold, width=2)
    draw.line([(width // 2 + 180, 443), (width // 2 + 250, 443)], fill=gold, width=2)
    
    # Vintage thin ornament/divider line under name
    draw.line([(width // 2 - 200, 485), (width // 2 + 200, 485)], fill=gold, width=1)
    draw.polygon([(width // 2, 480), (width // 2 + 8, 485), (width // 2, 490), (width // 2 - 8, 485)], fill=gold)
    
    # 11. Draw Description Text
    desc_text = "Edu Master Bot da o‘tkazilgan testda yuqori natija ko‘rsatgani uchun\nushbu sertifikat bilan taqdirlanadi."
    draw.text((width // 2, 535), desc_text, font=subtitle_font, fill=dark_gray, anchor="mm", align="center")
    
    # 12. Draw 5-Column Statistics with Circular Badges containing CUSTOM VECTOR ICONS (No emojis!)
    col_y = 650
    col_width_span = 960
    start_x = 160
    col_gap = col_width_span // 4
    
    for idx in range(5):
        cx = start_x + (idx * col_gap)
        
        # Draw dark navy circle badge
        draw.ellipse([cx - 28, col_y - 28, cx + 28, col_y + 28], fill=navy_blue, outline=gold, width=2)
        
        # Draw programmatic vector graphics inside circle depending on column
        if idx == 0:  # FAN (Book icon)
            draw.rectangle([cx - 12, col_y - 8, cx - 2, col_y + 8], fill="#FFFFFF")
            draw.rectangle([cx + 2, col_y - 8, cx + 12, col_y + 8], fill="#FFFFFF")
            draw.line([(cx - 9, col_y - 4), (cx - 5, col_y - 4)], fill=navy_blue, width=1)
            draw.line([(cx - 9, col_y), (cx - 5, col_y)], fill=navy_blue, width=1)
            draw.line([(cx - 9, col_y + 4), (cx - 5, col_y + 4)], fill=navy_blue, width=1)
            draw.line([(cx + 5, col_y - 4), (cx + 9, col_y - 4)], fill=navy_blue, width=1)
            draw.line([(cx + 5, col_y), (cx + 9, col_y)], fill=navy_blue, width=1)
            draw.line([(cx + 5, col_y + 4), (cx + 9, col_y + 4)], fill=navy_blue, width=1)
            title = "FAN"
            val = subject
        elif idx == 1:  # TEST KODI (Tag shape)
            draw.polygon([(cx - 12, col_y - 6), (cx + 4, col_y - 6), (cx + 12, col_y + 2), (cx + 4, col_y + 10), (cx - 12, col_y + 10)], fill="#FFFFFF")
            draw.ellipse([cx - 7, col_y - 1, cx - 3, col_y + 3], fill=navy_blue)
            title = "TEST KODI"
            val = test_code
        elif idx == 2:  # NATIJA (Award star medal)
            m_pts = []
            for i in range(10):
                angle = i * math.pi / 5 - math.pi / 2
                r = 12 if i % 2 == 0 else 6
                m_pts.append((cx + r * math.cos(angle), col_y - 2 + r * math.sin(angle)))
            draw.polygon(m_pts, fill=gold)
            draw.polygon([(cx - 5, col_y + 4), (cx - 9, col_y + 13), (cx - 1, col_y + 10)], fill="#FFFFFF")
            draw.polygon([(cx + 5, col_y + 4), (cx + 9, col_y + 13), (cx + 1, col_y + 10)], fill="#FFFFFF")
            title = "NATIJA"
            val = f"{percentage}%"
        elif idx == 3:  # TO'G'RI JAVOBLAR (Bullseye Target)
            draw.ellipse([cx - 12, col_y - 12, cx + 12, col_y + 12], outline="#FFFFFF", width=2)
            draw.ellipse([cx - 7, col_y - 7, cx + 7, col_y + 7], outline="#FFFFFF", width=1)
            draw.ellipse([cx - 3, col_y - 3, cx + 3, col_y + 3], fill=gold)
            title = "TO'G'RI JAVOBLAR"
            val = f"{correct} / {total}"
        else:  # SINF (Graduation Cap)
            draw.polygon([(cx, col_y - 8), (cx + 12, col_y - 2), (cx, col_y + 4), (cx - 12, col_y - 2)], fill="#FFFFFF")
            draw.chord([cx - 6, col_y, cx + 6, col_y + 10], 0, 180, fill="#FFFFFF")
            title = "SINF"
            val = class_name
            
        # Column Title (e.g. FAN)
        draw.text((cx, col_y + 48), title, font=sans_bold, fill=gold, anchor="mm")
        # Column Value (e.g. Matematika)
        draw.text((cx, col_y + 72), val, font=sans_bold, fill=navy_blue, anchor="mm")
        
    # 13. Bottom Section: Signature (Left), Seal/QR (Center), Date (Right)
    bottom_y = 835
    
    # A. Signature (Left) with elegant flowing calligraphy pen lines
    sig_points = []
    for t in range(50):
        x_sig = 120 + t * 4
        y_sig = bottom_y - 35 + 8 * math.sin(t * 0.3) - 4 * math.cos(t * 0.1)
        sig_points.append((x_sig, y_sig))
    draw.line(sig_points, fill=navy_blue, width=2)
    
    draw.line([(120, bottom_y - 10), (320, bottom_y - 10)], fill=dark_gray, width=1)
    draw.text((220, bottom_y + 10), "Bot Admini", font=sans_bold, fill=dark_gray, anchor="mm")
    draw.text((220, bottom_y + 28), "Edu Master Bot", font=caption_font, fill=navy_blue, anchor="mm")
    
    # B. Circular Stamp Logo (Middle-Left)
    stamp_cx, stamp_cy = 440, bottom_y - 10
    draw.ellipse([stamp_cx - 45, stamp_cy - 45, stamp_cx + 45, stamp_cy + 45], outline=navy_blue, width=2)
    draw.ellipse([stamp_cx - 40, stamp_cy - 40, stamp_cx + 40, stamp_cy + 40], outline=gold, width=1)
    draw.text((stamp_cx, stamp_cy - 12), "EDU MASTER", font=caption_font, fill=navy_blue, anchor="mm")
    
    # Draw stamp cap
    draw.polygon([(stamp_cx, stamp_cy - 2), (stamp_cx + 8, stamp_cy + 3), (stamp_cx, stamp_cy + 8), (stamp_cx - 8, stamp_cy + 3)], fill=navy_blue)
    draw.chord([stamp_cx - 4, stamp_cy + 4, stamp_cx + 4, stamp_cy + 12], 0, 180, fill=navy_blue)
    
    draw.text((stamp_cx, stamp_cy + 22), "KAFOLATLANGAN", font=caption_font, fill=navy_blue, anchor="mm")
    
    # C. QR Code inside a gold Wheat/Laurel Wreath border (Middle)
    qr_cx, qr_cy = 640, bottom_y - 15
    qr_data = f"Edu Master Bot\nSertifikat: {cert_id}\nIsm: {full_name}\nSinf: {class_name}\nFan: {subject}\nNatija: {percentage}%\nSana: {datetime.now().strftime('%d-%b-%Y')}"
    qr = qrcode.QRCode(version=1, box_size=3, border=1)
    qr.add_data(qr_data)
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color=navy_blue, back_color=background_color).convert("RGB")
    qr_img = qr_img.resize((90, 90))
    image.paste(qr_img, (qr_cx - 45, qr_cy - 45))
    
    # Programmatic golden Laurel Wreath arcs flanking the QR code
    draw.arc([qr_cx - 80, qr_cy - 45, qr_cx - 45, qr_cy + 45], 90, 270, fill=gold, width=2)
    draw.arc([qr_cx + 45, qr_cy - 45, qr_cx + 80, qr_cy + 45], 270, 90, fill=gold, width=2)
    for a in range(-40, 45, 20):
        lx = qr_cx - 62 + 8 * math.sin(a * math.pi / 180)
        ly = qr_cy + 35 * math.cos(a * math.pi / 180)
        draw.ellipse([lx - 3, ly - 3, lx + 3, ly + 3], fill=gold)
        
        rx = qr_cx + 62 - 8 * math.sin(a * math.pi / 180)
        ry = qr_cy + 35 * math.cos(a * math.pi / 180)
        draw.ellipse([rx - 3, ry - 3, rx + 3, ry + 3], fill=gold)
        
    draw.text((qr_cx, qr_cy + 65), f"Sertifikat ID: {cert_id}", font=sans_bold, fill=navy_blue, anchor="mm")
    
    # D. Calendar/Date (Right)
    sana_cx = 980
    # Programmatic grid calendar icon
    draw.rectangle([sana_cx - 16, bottom_y - 45, sana_cx + 16, bottom_y - 15], outline=navy_blue, width=2)
    draw.rectangle([sana_cx - 10, bottom_y - 49, sana_cx - 6, bottom_y - 43], fill=gold)
    draw.rectangle([sana_cx + 6, bottom_y - 49, sana_cx + 10, bottom_y - 43], fill=gold)
    draw.rectangle([sana_cx - 15, bottom_y - 44, sana_cx + 15, bottom_y - 36], fill=gold)
    for dx in [-8, 0, 8]:
        for dy in [-28, -20]:
            draw.rectangle([sana_cx + dx - 2, bottom_y + dy - 2, sana_cx + dx + 2, bottom_y + dy + 2], fill=navy_blue)
            
    draw.text((sana_cx, bottom_y + 5), "Sana", font=sans_bold, fill=gold, anchor="mm")
    current_date = datetime.now().strftime("%d-%b-%Y")
    draw.text((sana_cx, bottom_y + 25), current_date, font=sans_bold, fill=navy_blue, anchor="mm")
    
    # 14. Save Files securely
    Config.ensure_dirs()
    image_filename = f"cert_{cert_id.replace('-', '_')}.png"
    pdf_filename = f"cert_{cert_id.replace('-', '_')}.pdf"
    
    image_path = Config.CERTIFICATE_DIR / image_filename
    pdf_path = Config.CERTIFICATE_DIR / pdf_filename
    
    # Save Image
    image.save(image_path, "PNG")
    # Save PDF
    image.save(pdf_path, "PDF")
    
    return {
        "cert_id": cert_id,
        "image_path": str(image_path),
        "pdf_path": str(pdf_path)
    }


==================================================
FILE: .\utils\excel_export.py
==================================================

import os
import pandas as pd
from pathlib import Path
from datetime import datetime
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from config.config import Config

def export_test_results(results_data: list, test_code: str, subject: str, format_type: str) -> str:
    """
    Exports a list of result dictionaries to the requested format (excel, csv, pdf).
    Each dictionary in results_data must contain:
      - full_name, phone, class_name, telegram_id, test_code, subject, correct_count, wrong_count, percentage, submission_time
    Includes the 'Jami savollar' (Total count) column.
    Returns the absolute path to the generated file.
    """
    Config.ensure_dirs()
    
    # Add dynamic ranking
    sorted_results = sorted(
        results_data, 
        key=lambda x: (x.get("percentage", 0), -x.get("duration_spent", 999999)), 
        reverse=True
    )
    
    formatted_data = []
    for rank, r in enumerate(sorted_results, 1):
        formatted_data.append({
            "O'rin (Rank)": rank,
            "Ism va Familiya": r.get("full_name", ""),
            "Sinf": r.get("class_name", ""),
            "Test Kodi": r.get("test_code", test_code),
            "Fan": r.get("subject", subject),
            "Jami savollar": r.get("correct_count", 0) + r.get("wrong_count", 0),
            "To'g'ri javoblar": r.get("correct_count", 0),
            "Noto'g'ri javoblar": r.get("wrong_count", 0),
            "Natija (%)": f"{r.get('percentage', 0.0)}%",
            "Topshirilgan vaqt": r.get("submission_time", "")
        })

    df = pd.DataFrame(formatted_data)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = Config.BASE_DIR / "storage" / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    
    filename = f"natijalar_{test_code}_{timestamp}"
    
    if format_type.lower() == "excel":
        file_path = export_dir / f"{filename}.xlsx"
        
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        df.to_excel(writer, sheet_name='Natijalar', index=False)
        
        worksheet = writer.sheets['Natijalar']
        
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        navy_header_fill = PatternFill(start_color='0D2040', end_color='0D2040', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        
        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        alt_row_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')
        white_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        gold_row_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        
        data_font = Font(name='Arial', size=10)
        bold_font = Font(name='Arial', size=10, bold=True)
        
        for row_num in range(2, len(df) + 2):
            if row_num == 2:
                row_fill = gold_row_fill
                row_font = bold_font
            else:
                row_fill = alt_row_fill if (row_num % 2 == 0) else white_row_fill
                row_font = data_font
            
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = row_font
                
                col_name = df.columns[col_num - 1]
                if col_name in ["O'rin (Rank)", "Sinf", "Test Kodi", "Jami savollar", "To'g'ri javoblar", "Noto'g'ri javoblar", "Natija (%)", "Topshirilgan vaqt"]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        writer.close()
        return str(file_path)
        
    elif format_type.lower() == "csv":
        file_path = export_dir / f"{filename}.csv"
        df.to_csv(file_path, index=False, encoding="utf-8-sig")
        return str(file_path)
        
    elif format_type.lower() == "pdf":
        file_path = export_dir / f"{filename}.pdf"
        
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=landscape(letter),
            rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#0D2040'),
            alignment=1,
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            name='MetaStyle',
            fontName='Helvetica',
            fontSize=11,
            alignment=1,
            spaceAfter=20
        )
        
        story = []
        story.append(Paragraph(f"TEST NATIJALARI HISOBOTI", title_style))
        story.append(Paragraph(f"Fan: {subject} | Test kodi: {test_code} | Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        
        headers = ["O'rin", "Ism va Familiya", "Sinf", "Jami", "To'g'ri", "Noto'g'ri", "Natija", "Topshirilgan vaqt"]
        table_rows = [headers]
        
        for row in formatted_data:
            table_rows.append([
                str(row["O'rin (Rank)"]),
                row["Ism va Familiya"],
                row["Sinf"],
                str(row["Jami savollar"]),
                str(row["To'g'ri javoblar"]),
                str(row["Noto'g'ri javoblar"]),
                row["Natija (%)"],
                row["Topshirilgan vaqt"]
            ])
            
        col_widths = [40, 160, 50, 45, 50, 55, 52, 120]
        
        t = Table(table_rows, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0D2040')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('TOPPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#FFF2CC')),
            ('FONTNAME', (0,1), (-1,1), 'Helvetica-Bold'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0,2), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('FONTNAME', (0,2), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
        ]))
        
        story.append(t)
        doc.build(story)
        return str(file_path)
    else:
        raise ValueError(f"Unsupported format type: {format_type}")

def export_statistics_report(stats_data: dict, format_type: str) -> str:
    """
    Exports system statistics to Excel or PDF.
    """
    Config.ensure_dirs()
    
    formatted_data = [
        {"Ko'rsatkich (Metric)": "Jami ro'yxatdan o'tgan o'quvchilar", "Qiymat (Value)": f"{stats_data['total_users']} ta"},
        {"Ko'rsatkich (Metric)": "Aktiv (Test topshirgan) o'quvchilar", "Qiymat (Value)": f"{stats_data['active_users']} ta"},
        {"Ko'rsatkich (Metric)": "Jami topshirilgan testlar", "Qiymat (Value)": f"{stats_data['total_results']} ta"},
        {"Ko'rsatkich (Metric)": "Muvaffaqiyatli berilgan sertifikatlar", "Qiymat (Value)": f"{stats_data['certificates_count']} ta"},
        {"Ko'rsatkich (Metric)": "Tizimdagi eng yuqori natija", "Qiymat (Value)": f"{stats_data['highest_score']}%"},
        {"Ko'rsatkich (Metric)": "Tizimdagi eng past natija", "Qiymat (Value)": f"{stats_data['lowest_score']}%"},
        {"Ko'rsatkich (Metric)": "O'rtacha o'zlashtirish ko'rsatkichi", "Qiymat (Value)": f"{stats_data['average_score']}%"}
    ]
    
    df = pd.DataFrame(formatted_data)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = Config.BASE_DIR / "storage" / "exports"
    
    filename = f"tizim_statistikasi_{timestamp}"
    
    if format_type.lower() == "excel":
        file_path = export_dir / f"{filename}.xlsx"
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        df.to_excel(writer, sheet_name='Statistika', index=False)
        worksheet = writer.sheets['Statistika']
        
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        navy_header_fill = PatternFill(start_color='0D2040', end_color='0D2040', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        alt_row_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')
        white_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        
        for row_num in range(2, len(df) + 2):
            row_fill = alt_row_fill if (row_num % 2 == 0) else white_row_fill
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name='Arial', size=10)
                if col_num == 1:
                    cell.alignment = left_align
                else:
                    cell.alignment = center_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 6, 20)
            
        writer.close()
        return str(file_path)
        
    elif format_type.lower() == "pdf":
        file_path = export_dir / f"{filename}.pdf"
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=letter,
            rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#0D2040'),
            alignment=1,
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            name='MetaStyle',
            fontName='Helvetica',
            fontSize=11,
            alignment=1,
            spaceAfter=25
        )
        
        story = []
        story.append(Paragraph(f"TIZIM STATISTIKASI HISOBOTI", title_style))
        story.append(Paragraph(f"Eksport qilingan vaqt: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        
        headers = ["Ko'rsatkich (Metric)", "Qiymat (Value)"]
        table_rows = [headers]
        for row in formatted_data:
            table_rows.append([row["Ko'rsatkich (Metric)"], row["Qiymat (Value)"]])
            
        t = Table(table_rows, colWidths=[320, 180])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0D2040')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,0), 'CENTER'),
            ('ALIGN', (0,1), (0,-1), 'LEFT'),
            ('ALIGN', (1,1), (1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 11),
            ('BOTTOMPADDING', (0,0), (-1,0), 10),
            ('TOPPADDING', (0,0), (-1,0), 10),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 10),
            ('BOTTOMPADDING', (0,1), (-1,-1), 8),
            ('TOPPADDING', (0,1), (-1,-1), 8),
        ]))
        
        story.append(t)
        doc.build(story)
        return str(file_path)
    else:
        raise ValueError(f"Unsupported format type: {format_type}")

def export_users_list(users_data: list, format_type: str) -> str:
    """
    Exports all registered users list to Excel or PDF for administrator view.
    """
    Config.ensure_dirs()
    
    formatted_data = []
    for idx, u in enumerate(users_data, 1):
        formatted_data.append({
            "№": idx,
            "Ism va Familiya": u.get("full_name", ""),
            "Sinf": u.get("class_name", ""),
            "Telefon raqam": u.get("phone", ""),
            "Telegram ID": u.get("telegram_id", ""),
            "Ro'yxatdan o'tgan sana": u.get("registration_date", "")
        })
        
    df = pd.DataFrame(formatted_data)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = Config.BASE_DIR / "storage" / "exports"
    
    filename = f"foydalanuvchilar_ruyxati_{timestamp}"
    
    if format_type.lower() == "excel":
        file_path = export_dir / f"{filename}.xlsx"
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        df.to_excel(writer, sheet_name='Foydalanuvchilar', index=False)
        worksheet = writer.sheets['Foydalanuvchilar']
        
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        navy_header_fill = PatternFill(start_color='0D2040', end_color='0D2040', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        alt_row_fill = PatternFill(start_color='F5F7FA', end_color='F5F7FA', fill_type='solid')
        white_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        
        for row_num in range(2, len(df) + 2):
            row_fill = alt_row_fill if (row_num % 2 == 0) else white_row_fill
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name='Arial', size=10)
                
                col_name = df.columns[col_num - 1]
                if col_name in ["№", "Sinf", "Telegram ID", "Ro'yxatdan o'tgan sana"]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                    
        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        writer.close()
        return str(file_path)
        
    elif format_type.lower() == "pdf":
        file_path = export_dir / f"{filename}.pdf"
        doc = SimpleDocTemplate(
            str(file_path),
            pagesize=letter,
            rightMargin=30, leftMargin=30, topMargin=35, bottomMargin=30
        )
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name='TitleStyle',
            fontName='Helvetica-Bold',
            fontSize=15,
            textColor=colors.HexColor('#0D2040'),
            alignment=1,
            spaceAfter=15
        )
        
        meta_style = ParagraphStyle(
            name='MetaStyle',
            fontName='Helvetica',
            fontSize=10,
            alignment=1,
            spaceAfter=20
        )
        
        story = []
        story.append(Paragraph(f"RO'YXATDAN O'TGAN O'QUVCHILAR RO'YXATI", title_style))
        story.append(Paragraph(f"Jami ro'yxatdan o'tganlar: {len(users_data)} ta | Chop etilgan: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        
        headers = ["№", "Ism va Familiya", "Sinf", "Telefon raqam", "Telegram ID"]
        table_rows = [headers]
        
        for row in formatted_data:
            table_rows.append([
                str(row["№"]),
                row["Ism va Familiya"],
                row["Sinf"],
                row["Telefon raqam"],
                str(row["Telegram ID"])
            ])
            
        col_widths = [30, 200, 50, 110, 110]
        
        t = Table(table_rows, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0D2040')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('ALIGN', (1,1), (1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('TOPPADDING', (0,0), (-1,0), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F5F7FA')]),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
        ]))
        
        story.append(t)
        doc.build(story)
        return str(file_path)
    else:
        raise ValueError(f"Unsupported format type: {format_type}")


==================================================
FILE: .\utils\subscription.py
==================================================

import logging
from aiogram import Bot

logger = logging.getLogger("EduMasterBot.subscription")

async def is_subscribed(bot: Bot, channel_username: str, user_id: int) -> bool:
    # Ensure channel username starts with @ if it's not a numeric ID
    if not channel_username.startswith("@") and not channel_username.startswith("-"):
        channel_username = f"@{channel_username}"
        
    try:
        # Check if the channel is public and accessible
        member = await bot.get_chat_member(chat_id=channel_username, user_id=user_id)
        # Allowed states
        if member.status in ["creator", "administrator", "member"]:
            return True
    except Exception as e:
        logger.error(f"Error checking subscription for user {user_id} in {channel_username}: {e}")
        # In case the bot is not an administrator in the channel or cannot find it, return True to not lock out users,
        # or return False if strict subscription is required.
        # But wait, since mandatory subscription is specified: "MANDATORY CHANNEL SUBSCRIPTION"
        # We should assume strict checking. If there's an error (e.g. user not found), return False.
        # However, to be friendly during local tests if the channel doesn't exist, we can return True or False.
        # Let's return False for strict correctness, but we'll print a friendly instruction.
        return False
    return False
